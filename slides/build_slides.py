"""Builds the 17-slide deck (PPTX) following the suggested plan of the instruction sheet.
Rules of the instruction sheet: maximum 17 slides, headings 28-32 pt, body text at least 20 pt, own figures, presenter name in the corner."""
import os
import pickle

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG = os.path.join(ROOT, "figures")
RES = os.path.join(ROOT, "results")

NAVY = RGBColor(0x1F, 0x3A, 0x6E)
GREY = RGBColor(0x33, 0x33, 0x33)
ACC = RGBColor(0x1F, 0x78, 0xB4)
M = {"M1": "Aditya Khatri", "M2": "Simran Jeet Gill", "M3": "Divya Israni", "M4": "Amritanshi Singh", "M5": "Anchit Saxena", "All": "All members"}
MEMBERS = [("Aditya Khatri", "500119453"), ("Simran Jeet Gill", "500122195"), ("Divya Israni", "500121919"), ("Amritanshi Singh", "500119477"), ("Anchit Saxena", "500122868")]

df = pd.read_csv(os.path.join(RES, "all_runs.csv"))
df["succ"] = 100 * df.correct / df.n
CF = pd.read_csv(os.path.join(RES, "confirm_runs.csv"))
CF["succ"] = 100 * CF.correct / CF.n
CS = pd.read_csv(os.path.join(RES, "cs_runs.csv")) if os.path.exists(os.path.join(RES, "cs_runs.csv")) else None
if CS is not None:
    CS["succ"] = 100 * CS.correct / CS.n
sp = df[df.tag == "speed"].groupby(["method", "v"]).agg(succ=("succ", "mean"), thr=("throughput", "mean"), err=("grasp_err_mm", "mean"))
rip = df[df.tag == "ripple"].groupby(["method", "belt_ripple"]).agg(succ=("succ", "mean"))
M1, M2, M3 = "M1 Stop-and-go", "M2 Predictive", "M3 Proposed"
cab = {k: v.succ for k, v in CF[CF.tag == "confirm_ablation"].groupby("method")}
ab10 = df[df.tag == "ablation"].groupby("method").agg(succ=("succ", "mean"), err=("grasp_err_mm", "mean"))
with open(os.path.join(RES, "extra.pkl"), "rb") as _f:
    EX = pickle.load(_f)
TRK = {k: 1000 * float(abs(e).max()) for k, (t, e) in EX["track"].items()}
VIS = pd.read_csv(os.path.join(RES, "vision_study.csv"))
_v = VIS[(VIS.noise == 25) & (VIS.light == 1.0)].iloc[0]
V25, V25S = 100 * _v.color_correct / _v.objects, 100 * _v.shape_correct / _v.objects
_nom = df[~((df.tag == "noise") & (df.img_noise >= 60))]
FA = int(_nom.false_accept.sum() + CF.false_accept.sum())
FR = int(_nom.false_reject.sum() + CF.false_reject.sum())
NDEF = int(_nom.defects.sum() + CF.defects.sum())

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]
n_slide = [0]


def tb(slide, x, y, w, h, text, size=20, bold=False, color=GREY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = "Calibri"
    return box


def bullets(slide, x, y, w, h, items, size=20, gap=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, txt = (it if isinstance(it, tuple) else (None, it))
        r0 = p.add_run(); r0.text = "•  "; r0.font.size = Pt(size); r0.font.color.rgb = ACC; r0.font.bold = True
        if head:
            r1 = p.add_run(); r1.text = head + " "; r1.font.bold = True; r1.font.size = Pt(size); r1.font.color.rgb = NAVY
        r2 = p.add_run(); r2.text = txt; r2.font.size = Pt(size); r2.font.color.rgb = GREY
        for r in p.runs:
            r.font.name = "Calibri"
    return box


def new_slide(title, presenter):
    s = prs.slides.add_slide(blank)
    n_slide[0] += 1
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb(s, 0.5, 0.3, 9.4, 0.9, title, 30, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 9.6, 0.25, 3.5, 0.4, "Presenter: " + M[presenter], 16, True, ACC, PP_ALIGN.RIGHT)
    tb(s, 12.3, 7.05, 0.9, 0.35, str(n_slide[0]), 12, False, GREY, PP_ALIGN.RIGHT)
    tb(s, 0.5, 7.05, 8, 0.35, "CSAI4014 AISR | Project P14, Group G10, Batch B4", 12, False, GREY)
    return s


def pic(slide, name, x, y, w=None, h=None):
    path = os.path.join(FIG, name + ".png")
    if not os.path.exists(path):
        return
    im = Image.open(path)
    ar = im.width / im.height
    if w and h:
        if w / h > ar:
            w2 = h * ar; x += (w - w2) / 2; w = w2
        else:
            h2 = w / ar; y += (h - h2) / 2; h = h2
    elif w:
        h = w / ar
    else:
        w = h * ar
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def first(*names):
    for n in names:
        if os.path.exists(os.path.join(FIG, n + ".png")):
            return n
    return names[-1]


def table(slide, x, y, w, rows, col_w, size=20, row_h=0.6):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shape.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = str(val)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size); r.font.name = "Calibri"; r.font.bold = (i == 0)
                    r.font.color.rgb = RGBColor(255, 255, 255) if i == 0 else GREY
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xEE, 0xF3, 0xFA) if i % 2 else RGBColor(255, 255, 255))
    return t


DEMO = first("cs_still_descend", "still_descend")
SCENE = first("cs_still_transport", "fig_scene")

# 1 -----------------------------------------------------------------
s = new_slide("Vision-Guided Sorting Arm on a Conveyor", "M1")
tb(s, 0.5, 1.2, 12.3, 0.6, "Predictive visual servoing of a SCARA arm: colour and shape sorting in CoppeliaSim", 24, False, ACC)
tb(s, 0.5, 1.9, 12.3, 0.5, "Project P14 | Group G10 | Batch B4 | CSAI4014 Algorithms for Intelligent Systems and Robotics", 20, True, NAVY)
table(s, 0.5, 2.6, 6.8, [["Member", "SAP ID"]] + [list(m) for m in MEMBERS], (4.4, 2.4), 20, 0.6)
pic(s, SCENE, 7.6, 2.6, 5.4, 3.6)
tb(s, 0.5, 6.4, 12.3, 0.5, "Submitted to: Course Faculty, CSAI4014, School of Computer Science, UPES, Dehradun | October 2026", 20, False, GREY)

# 2 -----------------------------------------------------------------
s = new_slide("Problem and Motivation", "M1")
bullets(s, 0.5, 1.4, 6.3, 5.5, [
    ("Task:", "detect parts by colour AND shape on a moving belt; cubes to R / G / B bins, cylinders (wrong shape) to a reject bin"),
    ("Domain:", "quality inspection and sorting in manufacturing"),
    ("Hard part:", "the part moves during the 80 ms camera delay and the arm motion; belt speed ripples"),
    ("Today:", "many cells stop the belt and waste throughput"),
], 22)
pic(s, "fig_layout", 7.0, 1.4, 6.0, 4.6)

# 3 -----------------------------------------------------------------
s = new_slide("Objectives (measurable)", "M1")
table(s, 0.5, 1.4, 12.3, [
    ["#", "Objective", "Target", "Achieved"],
    ["O1", "Classify colour + shape", ">= 99 % at noise <= 25", "%.0f %% / %.0f %%" % (V25, V25S)],
    ["O2", "Grasp error, moving belt", "<= 5 mm to 0.35 m/s", "%.1f to %.1f mm" % (sp.loc[M3].err.min(), sp.loc[M3].err.max())],
    ["O3", "Sorting success, robust to ripple", ">= 95 % at <= 0.20 m/s", "%.1f %%; 15 %% ripple %.0f %%" % (sp.loc[(M3, 0.2)].succ, rip.loc[(M3, 0.15)].succ)],
    ["O4", "Throughput vs stop-and-go", ">= 2 x at 0.25 m/s", "%.1f x" % (sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr)],
    ["O5", "Inspection: wrong accept / reject", "<= 0.1 %", "%d of %d cylinders, %d rejects" % (FA, NDEF, FR)],
], (0.8, 4.6, 3.4, 3.5), 20, 0.78)
tb(s, 0.5, 6.2, 12.3, 0.8, "Assumptions: rendered orthographic camera, suction = attach constraint, kinematic conveyor, 3 colours, 2 shapes, simulation only.", 20, False, GREY)

# 4 -----------------------------------------------------------------
s = new_slide("Existing Solutions and the Gap", "M2")
table(s, 0.5, 1.4, 12.3, [
    ["Approach", "Strength", "Limitation"],
    ["Stop-and-go conveyor", "accurate, simple", "about half the throughput"],
    ["Encoder line tracking (FANUC, ABB)", "fast, mature", "closed software"],
    ["Visual servoing theory", "feedback cancels model error", "rarely applied to delayed moving parts"],
    ["Kalman tracking", "optimal estimate", "grasp benefit rarely quantified"],
], (4.6, 3.6, 4.1), 20, 0.8)
tb(s, 0.5, 5.8, 12.3, 1.0, "Gap: no open study that combines tracker + intercept planner + visual servo and compares it with baselines on two physics engines.", 22, True, NAVY)

# 5 -----------------------------------------------------------------
s = new_slide("Novelty", "M2")
bullets(s, 0.5, 1.4, 6.8, 5.6, [
    ("N1", "Visual-servo descent: grasp error %.1f to %.1f mm, success %.0f to %.0f %%" % (
        ab10.loc["M3 w/o visual servo"].err, ab10.loc[M3].err, ab10.loc["M3 w/o visual servo"].succ, ab10.loc[M3].succ)),
    ("N2", "Kalman + encoder: success %.0f to %.0f %% vs line fit (p < 0.001)" % (cab["M3 w/o Kalman (LS)"].mean(), cab[M3].mean())),
    ("N3", "Robust to 15 %% ripple: %.0f %% vs %.0f %%" % (rip.loc[(M3, 0.15)].succ, rip.loc[(M2, 0.15)].succ)),
    ("N4", "Same code on CoppeliaSim and MuJoCo, 890 + runs"),
    ("N5", "Negative result: scheduling rule gives no gain"),
], 20)
pic(s, "fig_ablation", 7.5, 1.7, 5.5, 3.3)

# 6 -----------------------------------------------------------------
s = new_slide("Syllabus Mapping", "M2")
table(s, 0.5, 1.4, 12.3, [
    ["AISR topic (Project List)", "Where it is used"],
    ["Robotic Perception", "vision sensor, HSV colour, contour shape, homography, Kalman tracking"],
    ["D-H Transformations", "D-H table, matrices A1 A2 A3, FK / IK, Jacobian"],
    ["Closed-Loop Control", "joint PID + feedforward, visual-servo descent"],
    ["Application Domains", "quality inspection and sorting in manufacturing"],
], (4.3, 8.0), 20, 0.95)

# 7 -----------------------------------------------------------------
s = new_slide("System Architecture", "M3")
pic(s, "fig_block_diagram", 0.5, 1.3, 12.3, 4.9)
tb(s, 0.5, 6.3, 12.3, 0.6, "20 fps vision sensor, 80 ms latency, encoder fusion, joint control inside the physics step", 20, True, NAVY)

# 8 -----------------------------------------------------------------
s = new_slide("Robot and Environment Model", "M3")
bullets(s, 0.5, 1.4, 5.8, 5.5, [
    ("Arm:", "SCARA, L1 0.35 m, L2 0.30 m, suction tool"),
    ("D-H:", "table + matrices, closed-form IK"),
    ("Physics:", "CoppeliaSim (Bullet), torque limits 45 / 18 N m"),
    ("Belt:", "0.05 to 0.35 m/s, 3 % ripple"),
    ("Parts:", "75 % cubes, 25 % cylinders"),
], 20)
pic(s, DEMO, 6.5, 1.5, 6.6, 3.9)
tb(s, 6.5, 5.5, 6.6, 0.6, "CoppeliaSim scene during the descent", 20, False, GREY)

# 9 -----------------------------------------------------------------
s = new_slide("Algorithm 1: Intercept Planning", "M3")
pic(s, "fig_flowchart", 0.4, 1.2, 6.2, 5.7)
bullets(s, 6.9, 1.5, 6.1, 5.3, [
    ("Predict:", "p(t) = p_hat + v_hat (t - t_f), compensates 80 ms"),
    ("Search:", "earliest t_a with IK solution and min-jerk time T <= available"),
    ("Select:", "shortest cycle among feasible targets"),
    ("Skip:", "parts that cannot be reached in time"),
], 20)

# 10 ----------------------------------------------------------------
s = new_slide("Algorithm 2 and Controller Design", "M4")
bullets(s, 0.5, 1.4, 5.6, 5.5, [
    ("Servo:", "v_cmd = Kt (p_hat(t_g) - p_tip), dq = J^-1 v_cmd"),
    ("PID:", "wn = 50 rad/s, zeta = 1, Ki = 0.2 Kp wn"),
    ("FF:", "J_n q_r'' + B_n q_r' cuts the PID tracking error by %.0f %%" % (100 * (1 - TRK["PID+FF"] / TRK["PID"]))),
    ("In CoppeliaSim:", "PID + FF runs in a joint-callback script"),
], 20)
pic(s, "fig_control", 6.2, 1.5, 6.9, 3.0)
tb(s, 6.2, 4.7, 6.9, 1.4, "Max tracking error (MuJoCo arm): PID %.1f mrad, PID + FF %.1f mrad" % (TRK["PID"], TRK["PID+FF"]), 20, False, GREY)

# 11 ----------------------------------------------------------------
s = new_slide("Simulation Setup", "M4")
bullets(s, 0.5, 1.4, 6.4, 5.5, [
    ("Platform:", "CoppeliaSim V4.10 (assigned) + Python + OpenCV"),
    ("Second engine:", "same code on MuJoCo for the large statistics"),
    ("Run:", "python demo_cs.py"),
    ("Check:", "success = part physically in its bin"),
], 20)
pic(s, "fig_camera", 7.1, 1.6, 5.9, 3.3)
tb(s, 7.1, 5.0, 5.9, 0.7, "Vision-sensor frame and detections", 20, False, GREY)

# 12 ----------------------------------------------------------------
s = new_slide("Live Demonstration", "M4")
tb(s, 0.5, 1.4, 6.7, 0.8, "Switch to CoppeliaSim now", 32, True, NAVY)
bullets(s, 0.5, 2.4, 6.8, 4.5, [
    ("Run:", "python demo_cs.py"),
    ("Show:", "arm waits ahead, descends while tracking; cubes to colour bins, cylinders to reject bin"),
    ("Compare:", "--method m2 --speed 0.3 shows misses"),
], 20)
pic(s, first("cs_still_transport", "still_transport"), 7.5, 2.4, 5.5, 3.3)

# 13 ----------------------------------------------------------------
s = new_slide("Results 1: Speed Study", "M5")
pic(s, "fig_speed", 0.4, 1.3, 12.5, 4.4)
tb(s, 0.5, 5.8, 12.3, 1.1, "M3: %.1f to %.1f mm grasp error, %.1f x the throughput of stop-and-go at 0.25 m/s. Above 0.25 m/s one arm cannot serve every part." % (
    sp.loc[M3].err.min(), sp.loc[M3].err.max(), sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr), 20, False, GREY)

# 14 ----------------------------------------------------------------
s = new_slide("Results 2: Comparison of Methods", "M5")
rows = [["Method (0.25 m/s)", "MuJoCo success %", "CoppeliaSim success %", "Grasp error mm"]]
for m_, lab in ((M1, "M1 stop-and-go"), (M2, "M2 predictive"), (M3, "M3 proposed")):
    cs = "-"
    if CS is not None:
        x = CS[(CS.tag == "speed") & (CS.method == m_) & (CS.v == 0.25)]
        if len(x):
            cs = "%.0f" % x.succ.mean()
    rows.append([lab, "%.0f" % sp.loc[(m_, 0.25)].succ, cs, "%.1f" % sp.loc[(m_, 0.25)].err])
table(s, 0.5, 1.4, 12.3, rows, (4.0, 2.8, 3.2, 2.3), 20, 0.7)
bullets(s, 0.5, 4.5, 12.3, 2.5, [
    ("Ripple 15 %:", "M3 %.0f %% vs M2 %.0f %% (30 seeds: p < 0.001)" % (rip.loc[(M3, 0.15)].succ, rip.loc[(M2, 0.15)].succ)),
    ("Ablation:", "M2 %.0f, no servo %.0f, no Kalman %.0f, full %.0f %%" % (cab[M2].mean(), cab["M3 w/o visual servo"].mean(), cab["M3 w/o Kalman (LS)"].mean(), cab[M3].mean())),
    ("Both engines:", "same ranking M3 > M2"),
], 20, 4)

# 15 ----------------------------------------------------------------
s = new_slide("Team Contribution", "M5")
table(s, 0.5, 1.4, 12.3, [
    ["Member", "Technical module"],
    ["Aditya Khatri", "SCARA kinematics, D-H, CoppeliaSim and MuJoCo cells, visual servo"],
    ["Simran Jeet Gill", "Vision sensor image, HSV detector, shape classifier, homography"],
    ["Divya Israni", "Kalman tracker, intercept planner, ablation"],
    ["Amritanshi Singh", "PID + feedforward, joint callback, tuning, step tests"],
    ["Anchit Saxena", "Experiments, stop-and-go baseline, renders, live demo"],
], (3.2, 9.1), 20, 0.82)

# 16 ----------------------------------------------------------------
s = new_slide("Conclusion and Industry Use", "All")
bullets(s, 0.5, 1.4, 7.2, 5.5, [
    ("Findings:", "visual servo and Kalman both help; scheduling did not"),
    ("Domain:", "quality inspection and sorting in manufacturing"),
    ("Real robot:", "triggered camera, hand-eye calibration, vacuum delay"),
    ("Future:", "second arm, learned detector, ROS and real-arm test"),
    ("Team:", "we taught each other our modules"),
], 20)
pic(s, "fig_trace", 7.9, 1.8, 5.2, 2.6)
tb(s, 7.9, 4.5, 5.2, 0.9, "Same trial: M2 has misses above 10 mm, M3 none", 20, False, GREY)

# 17 ----------------------------------------------------------------
s = new_slide("Thank You", "All")
tb(s, 0.5, 1.8, 12.3, 1.0, "Questions?", 40, True, NAVY)
tb(s, 0.5, 3.2, 12.3, 0.6, "GitHub: https://github.com/Aditya-Khatri-7/Robotics-assignment-project-b4", 22, False, GREY)
tb(s, 0.5, 3.9, 12.3, 1.0, "E-mails (@stu.upes.ac.in): aditya.119453, simran.122195, divya.121919, amritanshi.119477, anchit.122868", 20, False, GREY)
tb(s, 0.5, 5.0, 12.3, 0.9, "Group G10 | Aditya Khatri, Simran Jeet Gill, Divya Israni, Amritanshi Singh, Anchit Saxena", 20, False, GREY)

NOTES = [
    "Introduce yourself with name and SAP ID. Project P14: vision-guided sorting arm on a conveyor, built in CoppeliaSim (the assigned platform) with Python and OpenCV and cross-checked in MuJoCo.",
    "Task: the camera must decide colour AND shape. Cubes go to the R/G/B bins, cylinders (wrong shape) go to the reject bin. Domain: quality inspection and sorting in manufacturing. Say why stopping the belt wastes throughput.",
    "Read the Achieved column. Five objectives. State the assumptions honestly.",
    "Stop-and-go is accurate but slow, commercial line tracking is closed, visual servoing and Kalman tracking are known but rarely compared on moving, delayed targets. That is our gap.",
    "Five novelty points with numbers. Say the scheduling rule gave no gain: an honest negative result. Mention the 30-seed confirmation study.",
    "Use the official four topics of the Project List: Robotic Perception, D-H Transformations, Closed-Loop Control, Application Domains.",
    "Follow one part from the vision sensor to the bin. Explain encoder fusion and the 80 ms latency.",
    "Arm L1 0.35 m, L2 0.30 m, rigid-body physics with torque limits. Parts: 75 percent cubes, 25 percent cylinders. CoppeliaSim scene on the right.",
    "Explain the flowchart: predict, find the earliest feasible intercept time with inverse kinematics and minimum-jerk time, pick the shortest cycle. Parts that cannot be reached in time are skipped.",
    "The visual servo aims at the predicted contact point, not the present position. Gains from natural frequency 50 rad/s and damping 1. In CoppeliaSim the PID runs in a joint callback script.",
    "Platform and commands. Show the vision-sensor frame with colour and shape labels. Mention the physical sorting check and the unit tests.",
    "Run python demo_cs.py live. Point out: arm waits ahead of the part, descends while correcting, cubes into colour bins, cylinders into the reject bin. Then run --method m2 --speed 0.3 to show misses.",
    "Stop-and-go is 100 percent but slow. M3 keeps about 3 mm error. Above 0.25 m/s one arm cannot serve all parts: capacity limit, not accuracy.",
    "Table compares methods on both engines. Ripple, ablation and p-values from the confirmation study. Same ranking on CoppeliaSim and MuJoCo.",
    "Each member states their own technical module and the files they wrote. Every claim is traceable to commits and to the video.",
    "Findings, industry use, transfer to a real robot, future work. Say what the team learned together.",
    "Thank the audience, show the GitHub link and take questions. Every member should be ready to answer about the whole project.",
]
for _sl, _nt in zip(prs.slides, NOTES):
    _sl.notes_slide.notes_text_frame.text = _nt
assert len(prs.slides) == 17
out = os.path.join(ROOT, "AISR_G10_Slides.pptx")
prs.save(out)
print("saved", out)
