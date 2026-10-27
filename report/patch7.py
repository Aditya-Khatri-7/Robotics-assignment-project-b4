import os, re, shutil
R = 'D:/!aRobotics project/sortarm_project/'
os.makedirs(R + 'report/template', exist_ok=True)
shutil.copy('D:/!aRobotics project/provided_zip/b3+b4/AISR_Sample_Project_Report.docx', R + 'report/template/AISR_Sample_Project_Report.docx')
p = R + 'report/build_report.py'
s = open(p, encoding='utf8').read()


def rep(a, b):
    global s
    assert a in s, a[:90]
    s = s.replace(a, b, 1)


# ------------------------------------------------ class setup on the sample template
a = s.index("class B:")
b_ = s.index("    # ---------------------------------------------------------------- text helpers")
s = s[:a] + '''SAMPLE = os.path.join(HERE, "template", "AISR_Sample_Project_Report.docx")


class B:
    def __init__(self):
        self.d = Document(SAMPLE)
        self.cites = []
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self.sec_pages = {}
        self._trim()
        self._setup()

    def _trim(self):
        """keep only the official title page (with its section break), the front-matter section break and the final section"""
        body = self.d.element.body
        kids = list(body)
        keep = set(range(0, 18)) | {56, len(kids) - 1}
        for i, el in enumerate(kids):
            if i not in keep:
                body.remove(el)
        self.sp1 = kids[56]

    def begin_front(self):
        pass

    def begin_body(self):
        """front-matter elements were appended after the section break paragraph: move them in front of it"""
        body = self.d.element.body
        kids = list(body)
        i1 = kids.index(self.sp1)
        for el in [e for e in kids[i1 + 1:] if e.tag != qn("w:sectPr")]:
            self.sp1.addprevious(el)

    # ---------------------------------------------------------------- setup
    def _setup(self):
        st = self.d.styles
        from docx.enum.style import WD_STYLE_TYPE
        set_font(st["Normal"], 12)
        for name, size in (("Heading 1", 14), ("Heading 2", 12)):
            set_font(st[name], size, True, False)
        set_font(st["Caption"], 11, False, False)
        fmt(st["Caption"], WD_ALIGN_PARAGRAPH.CENTER, 4, 10, 1.0)
        for nm in ("List Bullet", "List Number"):
            try:
                set_font(st[nm], 12)
                fmt(st[nm], WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, 1.5)
            except KeyError:
                pass
        for nm, ind in (("TOC 1", 0), ("TOC 2", 0.6), ("Table of Figures", 0)):
            try:
                sty = st[nm]
            except KeyError:
                sty = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
                sty.base_style = st["Normal"]
            set_font(sty, 10.5)
            fmt(sty, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0)
            sty.paragraph_format.left_indent = Cm(ind)

''' + s[b_:]

# ------------------------------------------------ captions with colon like the sample
rep('        self._runs(c, ". " + self._tok(caption))\n\n    def tab_caption', '        self._runs(c, ": " + self._tok(caption))\n\n    def tab_caption')
rep('        c.add_run("Table ")\n        add_field(c, "SEQ Table \\\\* ARABIC", str(self.tab_n))\n        self._runs(c, ". " + self._tok(caption))',
    '        c.add_run("Table ")\n        add_field(c, "SEQ Table \\\\* ARABIC", str(self.tab_n))\n        self._runs(c, ": " + self._tok(caption))')

# ------------------------------------------------ headings numbered like the sample ("1.  Introduction")
rep('''    def h1(self, text, page_break=False):
        p = self.d.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = page_break
        self._runs(p, text)
        return p''', '''    def h1(self, text, page_break=False):
        p = self.d.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = page_break
        text = re.sub(r"^13 References$", "References", text)
        text = re.sub(r"^(\\d+) ", r"\\1.  ", text)
        self._runs(p, text)
        return p''')

# ------------------------------------------------ title page: fill the sample's own title page
a = s.index("    # ============================================================ TITLE PAGE")
b_ = s.index("    # ============================================================ FRONT MATTER (roman)")
s = s[:a] + '''    # ============================================================ TITLE PAGE (official sample, only the variable parts change)
    def set_text(par, new):
        runs = par.runs
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""
    paras = d.paragraphs
    set_text(paras[3], TITLE)
    for par in paras[7:10]:
        for r in par.runs:
            r.text = r.text.replace("P00 (Sample)", "P14").replace("G00", "G10").replace("B1", "B4")
    tt = d.tables[0]
    for i, (n, sid) in enumerate(MEMBERS, 1):
        for j, v in enumerate((str(i), n, sid)):
            cell = tt.cell(i, j)
            par = cell.paragraphs[0]
            if par.runs:
                par.runs[0].text = v
                for r in par.runs[1:]:
                    r.text = ""
            else:
                par.add_run(v)

''' + s[b_:]

# ------------------------------------------------ front matter
rep('    b.new_section("lowerRoman", 1)\n    b.h1("Declaration")', '    b.begin_front()\n    b.h1("Declaration")')
rep('''    par = d.add_paragraph(); par.paragraph_format.space_before = Pt(16)
    par.add_run("Date: ____ / ____ / 2026")''', '''    par = d.add_paragraph(); par.paragraph_format.space_before = Pt(16); par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.add_run("Date: ____________________\\tPlace: Dehradun")''')
rep('    b.h1("Index (Table of Contents)", True)', '    b.h1("Index", True)')
rep('    b.new_section("decimal", 1)\n', '    b.begin_body()\n')
# abstract keywords (after the abstract paragraph)
rep('''            sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr))

    b.h1("Index", True)''', '''            sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr))
    kw = d.add_paragraph()
    kw.add_run("Keywords: ").bold = True
    kw.add_run("SCARA arm, conveyor sorting, visual servoing, Kalman filter, D-H transformations, MuJoCo, OpenCV.")

    b.h1("Index", True)''')
# writing allocation table like the sample
a = s.index('    alloc = [("Aditya Khatri (500119453)"')
b_ = s.index('    b.table("alloc", "Writing allocation of the report.", rows, (5.0, 8.0, 2.5))')
s = s[:a] + '''    alloc = [("Aditya Khatri (500119453)", "Abstract, 1, 2, 12"),
             ("Simran Jeet Gill (500122195)", "3, 4, 8, References"),
             ("Divya Israni (500121919)", "5, 6"),
             ("Amritanshi Singh (500119477)", "7, 9.1-9.2, Appendix"),
             ("Anchit Saxena (500122868)", "9.3-9.7")]
    rows = [["Member", "Name and SAP ID", "Sections written", "Pages"]]
    for i, (n, secs) in enumerate(alloc):
        rows.append(["M%d" % (i + 1), n, secs, pages.get("alloc%d" % i, "")])
    rows.append(["All", "Every member", "10 (own row and signature), 11 (own reflection), Declaration", "-"])
''' + s[b_:]
rep('    b.table("alloc", "Writing allocation of the report.", rows, (5.0, 8.0, 2.5))', '    b.table("alloc", "Writing allocation of the report.", rows, (1.6, 5.0, 6.9, 2.0))')
# contribution sections naming
s = s.replace('"3, 4, 8, 13"', '"3, 4, 8, References"')
# meeting log with No. column
rep('    rows = [["Date", "Duration", "Attendees", "Decisions"]]', '    rows = [["No.", "Date", "Duration", "Attendees", "Decisions"]]')
rep('        rows.append(["__/__/2026", "__ min", "all five members", dec])', '        rows.append([str(len(rows)), "__/__/2026", "__ min", "All 5", dec])')
rep('    b.table("log", "Meeting log of the group.", rows, (2.6, 2.0, 3.2, 7.7), size=9.5)', '    b.table("log", "Meeting log of the group.", rows, (1.0, 2.5, 1.9, 1.8, 8.3), size=9.5)')
open(p, 'w', encoding='utf8').write(s)

# ------------------------------------------------ measure_pages headings
m = R + 'report/measure_pages.py'
t = open(m, encoding='utf8').read()
t = t.replace('"13 References"', '"References"')
t = re.sub(r'"(\d+) ([A-Z][^"]*)"', lambda mo: '"%s.  %s"' % (mo.group(1), mo.group(2)) if '.' not in mo.group(1) else mo.group(0), t)
t = t.replace('"9.3.  Comparison of the three strategies"', '"9.3 Comparison of the three strategies"')
t = t.replace('b[4].strip().split("\\n")[0].strip() == h', '" ".join(b[4].strip().split("\\n")[0].split()) == " ".join(h.split())')
open(m, 'w', encoding='utf8').write(t)
print('ok')
