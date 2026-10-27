"""Builds the AISR project report (docx) in the format required by the instruction sheet."""
import json
import os
import re
import sys

import pandas as pd
import pickle
import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, ROOT)
from docx_lib import *  # noqa
import mujoco
from dataclasses import replace
from sortarm.config import Config
from sortarm.mj_world import World
_w = World(Config(), 1, ['red'])
J_NOM = tuple(float(v) for v in _w.nominal_inertia((0.3, 1.5)))

FIG = os.path.join(ROOT, "figures")
RES = os.path.join(ROOT, "results")

TITLE = "Vision-Guided Sorting Arm on a Conveyor"
SUBTITLE = "Predictive Visual Servoing of a SCARA Arm for Colour Sorting of Moving Objects"
MEMBERS = [("Aditya Khatri", "500119453"), ("Simran Jeet Gill", "500122195"), ("Divya Israni", "500121919"),
           ("Amritanshi Singh", "500119477"), ("Anchit Saxena", "500122868")]
GH = "https://github.com/Aditya-Khatri-7/Robotics-assignment-project-b4"
YT = "https://youtu.be/<unlisted-video-id>"

REFS = {
    "sic": 'B. Siciliano, L. Sciavicco, L. Villani, and G. Oriolo, <i>Robotics: Modelling, Planning and Control</i>. London, U.K.: Springer, 2009.',
    "kal": 'R. E. Kalman, "A new approach to linear filtering and prediction problems," <i>J. Basic Eng.</i>, vol. 82, no. 1, pp. 35-45, Mar. 1960.',
    "chau1": 'F. Chaumette and S. Hutchinson, "Visual servo control, Part I: Basic approaches," <i>IEEE Robot. Autom. Mag.</i>, vol. 13, no. 4, pp. 82-90, Dec. 2006.',
    "chau2": 'F. Chaumette and S. Hutchinson, "Visual servo control, Part II: Advanced approaches," <i>IEEE Robot. Autom. Mag.</i>, vol. 14, no. 1, pp. 109-118, Mar. 2007.',
    "corke": 'P. Corke, <i>Robotics, Vision and Control: Fundamental Algorithms in MATLAB</i>, 2nd ed. Cham, Switzerland: Springer, 2017.',
    "flash": 'T. Flash and N. Hogan, "The coordination of arm movements: An experimentally confirmed mathematical model," <i>J. Neurosci.</i>, vol. 5, no. 7, pp. 1688-1703, Jul. 1985.',
    "cv": 'G. Bradski, "The OpenCV library," <i>Dr. Dobb\'s J. Softw. Tools</i>, vol. 25, no. 11, pp. 120-125, Nov. 2000.',
    "hz": 'R. Hartley and A. Zisserman, <i>Multiple View Geometry in Computer Vision</i>, 2nd ed. Cambridge, U.K.: Cambridge Univ. Press, 2004.',
    "craig": 'J. J. Craig, <i>Introduction to Robotics: Mechanics and Control</i>, 4th ed. New York, NY, USA: Pearson, 2018.',
    "astrom": 'K. J. Åström and T. Hägglund, <i>PID Controllers: Theory, Design, and Tuning</i>, 2nd ed. Research Triangle Park, NC, USA: ISA, 1995.',
    "bs": 'Y. Bar-Shalom, X. R. Li, and T. Kirubarajan, <i>Estimation with Applications to Tracking and Navigation</i>. New York, NY, USA: Wiley, 2001.',
    "mujoco": 'E. Todorov, T. Erez, and Y. Tassa, "MuJoCo: A physics engine for model-based control," in <i>Proc. IEEE/RSJ Int. Conf. Intell. Robots Syst. (IROS)</i>, Vilamoura, Portugal, Oct. 2012, pp. 5026-5033.',
    "coppelia": 'E. Rohmer, S. P. N. Singh, and M. Freese, "CoppeliaSim (formerly V-REP): A versatile and scalable robot simulation framework," in <i>Proc. IEEE/RSJ Int. Conf. Intell. Robots Syst. (IROS)</i>, Tokyo, Japan, Nov. 2013, pp. 1321-1326.',
    "gw": 'R. C. Gonzalez and R. E. Woods, <i>Digital Image Processing</i>, 4th ed. New York, NY, USA: Pearson, 2018.',
    "dhorig": 'J. Denavit and R. S. Hartenberg, "A kinematic notation for lower-pair mechanisms based on matrices," <i>J. Appl. Mech.</i>, vol. 22, no. 2, pp. 215-221, Jun. 1955.',
    "iso": '<i>Manipulating Industrial Robots - Performance Criteria and Related Test Methods</i>, ISO Standard 9283, 1998.',
    "clavel": 'R. Clavel, "Delta, a fast robot with parallel geometry," in <i>Proc. 18th Int. Symp. Ind. Robots (ISIR)</i>, Sydney, Australia, Apr. 1988, pp. 91-100.',
    "fanuc": 'FANUC Corporation, "iRVision integrated robot vision and line tracking," product documentation, Oshino, Japan. [Online]. Available: https://www.fanuc.eu. <hl>(accessed: 27 Oct. 2026)</hl>',
    "abb": 'ABB Robotics, "PickMaster conveyor-tracking software for FlexPicker robots," product information, Zürich, Switzerland. [Online]. Available: https://new.abb.com/products/robotics. <hl>(accessed: 27 Oct. 2026)</hl>',
}

FIG_ORDER = ["layout", "block", "scene", "flow", "camera", "control", "vision", "speed", "stress", "trace", "snap", "ablation"]
FIG_FILE = {"layout": "fig_layout", "block": "fig_block_diagram", "scene": "fig_scene", "camera": "fig_camera", "flow": "fig_flowchart",
            "control": "fig_control", "vision": "fig_vision", "speed": "fig_speed", "stress": "fig_stress",
            "trace": "fig_trace", "ablation": "fig_ablation", "snap": "fig_snapshots"}
TAB_ORDER = ["alloc", "obj", "lit", "nov", "syl", "dh", "par", "meth", "map", "sim", "val", "speed", "abl", "conf", "conf2", "cs", "contrib", "log", "peer"]


SAMPLE = os.path.join(HERE, "template", "AISR_Sample_Project_Report.docx")


class B:
    def __init__(self):
        self.d = Document(SAMPLE)
        self.cites = []
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self.sec_pages = {}
        self._trim()
        self._setup()

    def _trim(self):
        """keep only the official title page (with its section break), the front-matter section break and the final section"""
        body = self.d.element.body
        kids = list(body)
        keep = set(range(0, 18)) | {56, len(kids) - 1}
        for i, el in enumerate(kids):
            if i not in keep:
                body.remove(el)
        self.sp1 = kids[56]

    def begin_front(self):
        pass

    def begin_body(self):
        """front-matter elements were appended after the section break paragraph: move them in front of it"""
        body = self.d.element.body
        kids = list(body)
        i1 = kids.index(self.sp1)
        for el in [e for e in kids[i1 + 1:] if e.tag != qn("w:sectPr")]:
            self.sp1.addprevious(el)

    # ---------------------------------------------------------------- setup
    def _setup(self):
        st = self.d.styles
        from docx.enum.style import WD_STYLE_TYPE
        set_font(st["Normal"], 12)
        for name, size in (("Heading 1", 14), ("Heading 2", 12)):
            set_font(st[name], size, True, False)
        set_font(st["Caption"], 11, False, False)
        fmt(st["Caption"], WD_ALIGN_PARAGRAPH.CENTER, 4, 10, 1.0)
        for nm in ("List Bullet", "List Number"):
            try:
                set_font(st[nm], 12)
                fmt(st[nm], WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, 1.5)
            except KeyError:
                pass
        for nm, ind in (("TOC 1", 0), ("TOC 2", 0.6), ("Table of Figures", 0)):
            try:
                sty = st[nm]
            except KeyError:
                sty = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
                sty.base_style = st["Normal"]
            set_font(sty, 10.5)
            fmt(sty, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0)
            sty.paragraph_format.left_indent = Cm(ind)

    # ---------------------------------------------------------------- text helpers
    def _tok(self, text):
        def fig(m):
            return "Figure %d" % (FIG_ORDER.index(m.group(1)) + 1)

        def tab(m):
            return "Table %d" % (TAB_ORDER.index(m.group(1)) + 1)

        def cite(m):
            out = []
            for k in m.group(1).split(","):
                k = k.strip()
                if k not in self.cites:
                    self.cites.append(k)
                out.append("[%d]" % (self.cites.index(k) + 1))
            return ", ".join(out)
        text = re.sub(r"\{F:(\w+)\}", fig, text)
        text = re.sub(r"\{T:(\w+)\}", tab, text)
        text = re.sub(r"\{c:([\w,\s]+)\}", cite, text)
        return text

    def _runs(self, par, text, size=None, base_bold=None):
        """Inline markup: **bold**, <i>italic</i>, <hl>highlight</hl>."""
        parts = re.split(r"(\*\*.+?\*\*|<i>.+?</i>|<hl>.+?</hl>)", text)
        for p in parts:
            if not p:
                continue
            if p.startswith("**"):
                r = par.add_run(p[2:-2]); r.bold = True
            elif p.startswith("<i>"):
                r = par.add_run(p[3:-4]); r.italic = True
            elif p.startswith("<hl>"):
                r = par.add_run(p[4:-5]); highlight(r)
            else:
                r = par.add_run(p)
                if base_bold:
                    r.bold = True
            if size:
                r.font.size = Pt(size)

    def h1(self, text, page_break=False):
        p = self.d.add_paragraph(style="Heading 1")
        p.paragraph_format.page_break_before = page_break
        text = re.sub(r"^13 References$", "References", text)
        text = re.sub(r"^(\d+) ", r"\1.  ", text)
        self._runs(p, text)
        return p

    def h2(self, text):
        p = self.d.add_paragraph(style="Heading 2")
        self._runs(p, text)
        return p

    def p(self, text, **kw):
        par = self.d.add_paragraph()
        self._runs(par, self._tok(text))
        return par

    def bullets(self, items, style="List Bullet"):
        for t in items:
            par = self.d.add_paragraph(style=style)
            self._runs(par, self._tok(t))

    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def eq(self, text):
        self.eq_n += 1
        p = self.d.add_paragraph()
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.15
        pf.space_before = Pt(2); pf.space_after = Pt(4)
        pf.tab_stops.add_tab_stop(Cm(7.75), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t" + text + "\t(%d)" % self.eq_n)
        r.italic = True
        r.font.size = Pt(11.5)
        return p

    def fig(self, key, caption, width_cm=14.0):
        self.fig_n += 1
        assert FIG_ORDER[self.fig_n - 1] == key, (key, self.fig_n)
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(os.path.join(FIG, FIG_FILE[key] + ".png"), width=Cm(width_cm))
        c = self.d.add_paragraph(style="Caption")
        c.add_run("Figure ")
        add_field(c, "SEQ Figure \\* ARABIC", str(self.fig_n))
        self._runs(c, ": " + self._tok(caption))

    def tab_caption(self, key, caption):
        self.tab_n += 1
        assert TAB_ORDER[self.tab_n - 1] == key, (key, self.tab_n)
        c = self.d.add_paragraph(style="Caption")
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_after = Pt(4)
        c.add_run("Table ")
        add_field(c, "SEQ Table \\* ARABIC", str(self.tab_n))
        self._runs(c, ": " + self._tok(caption))

    def table(self, key, caption, rows, widths, size=10, header=True, align_left=True, keep=True):
        self.tab_caption(key, caption)
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths[j])
                par = cell.paragraphs[0]
                par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT if align_left else WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(2); par.paragraph_format.space_before = Pt(2)
                self._runs(par, self._tok(str(val)), size=size, base_bold=(header and i == 0))
                if header and i == 0:
                    shade(cell)
        for row in (t.rows[:-1] if keep else t.rows[:1]):
            for cell in row.cells:
                for pp in cell.paragraphs:
                    pp.paragraph_format.keep_with_next = True
        sp = self.d.add_paragraph()
        sp.paragraph_format.space_after = Pt(4); sp.paragraph_format.line_spacing = 1.0
        return t

    def code(self, lines, size=8.5):
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.width = Cm(15.0)
        shade(cell, "F2F2F2")
        first = True
        for ln in lines:
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.keep_together = True
            par.paragraph_format.keep_with_next = ln is not lines[-1]
            r = par.add_run(ln)
            r.font.name = "Courier New"
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Courier New")
            r.font.size = Pt(size)
        sp = self.d.add_paragraph(); sp.paragraph_format.space_after = Pt(2); sp.paragraph_format.line_spacing = 1.0

    def new_section(self, fmt_, start, numbered=True):
        s = self.d.add_section(WD_SECTION.NEW_PAGE)
        page_border(s)
        page_numbering(s, fmt_, start)
        if numbered:
            footer_pagenum(s)
        else:
            footer_empty(s)
        return s


def load_data():
    df = pd.read_csv(os.path.join(RES, "all_runs.csv"))
    df["succ"] = 100 * df.correct / df.n
    with open(os.path.join(RES, "extra.pkl"), "rb") as f:
        ex = pickle.load(f)
    return df, ex


def build(out_path, pages=None):
    df, ex = load_data()
    b = B()
    d = b.d
    pages = pages or {}

    # numbers used in the text -------------------------------------------------
    sp = df[df.tag == "speed"].groupby(["method", "v"]).agg(succ=("succ", "mean"), thr=("throughput", "mean"),
                                                            err=("grasp_err_mm", "mean"), sd=("succ", "std"))
    ab = df[df.tag == "ablation"].groupby("method").agg(succ=("succ", "mean"), sd=("succ", "std"), thr=("throughput", "mean"),
                                                         err=("grasp_err_mm", "mean"))
    lat = df[df.tag == "latency"].groupby(["method", "latency"]).agg(succ=("succ", "mean"), err=("grasp_err_mm", "mean"))
    rip = df[df.tag == "ripple"].groupby(["method", "belt_ripple"]).agg(succ=("succ", "mean"), err=("grasp_err_mm", "mean"))
    noi = df[df.tag == "noise"].groupby(["method", "img_noise"]).agg(succ=("succ", "mean"), err=("grasp_err_mm", "mean"))
    M1, M2, M3 = "M1 Stop-and-go", "M2 Predictive", "M3 Proposed"
    vis = pd.read_csv(os.path.join(RES, "vision_study.csv"))
    vis["acc"] = 100 * vis.color_correct / vis.objects
    vis["acc_shape"] = 100 * vis.shape_correct / vis.objects
    from scipy import stats as _st

    def welch(a, b):
        return float(_st.ttest_ind(a, b, equal_var=False).pvalue)

    def pfmt(pv):
        return "< 0.001" if pv < 0.001 else "%.3f" % pv
    cs_ = pd.read_csv(os.path.join(RES, "cs_runs.csv"))
    cs_["succ"] = 100 * cs_.correct / cs_.n
    csx = {}
    cs_ctl = None
    if os.path.exists(os.path.join(RES, "cs_control.pkl")):
        with open(os.path.join(RES, "cs_control.pkl"), "rb") as _f:
            cs_ctl = pickle.load(_f)
    cf = pd.read_csv(os.path.join(RES, "confirm_runs.csv"))
    cf["succ"] = 100 * cf.correct / cf.n
    cab = {m_: x_ for m_, x_ in cf[cf.tag == "confirm_ablation"].groupby("method")}
    hi = (df.tag == "noise") & (df.img_noise >= 60)            # extreme image noise (7.5 x the default)
    nom = df[~hi]
    n_def, n_all = int(nom.defects.sum()), int(nom.n.sum())
    n_fa, n_fr, n_ce = int(nom.false_accept.sum()), int(nom.false_reject.sum()), int(nom.colour_err.sum())
    n_good = n_all - n_def
    h_def, h_fa, h_n = int(df[hi].defects.sum()), int(df[hi].false_accept.sum()), int(df[hi].n.sum())
    c_def, c_fa, c_fr, c_n = int(cf.defects.sum()), int(cf.false_accept.sum()), int(cf.false_reject.sum()), int(cf.n.sum())
    a_def, a_fa, a_fr, a_good = n_def + c_def, n_fa + c_fa, n_fr + c_fr, (n_all - n_def) + (c_n - c_def)
    step = {}
    for k, (t, y) in ex["step"].items():
        idx = np.where(abs(y - 0.5) > 0.01)[0]
        step[k] = (100 * (y.max() - 0.5) / 0.5, float(t[idx[-1]]) if len(idx) else 0.0)
    trk = {k: 1000 * float(abs(e).max()) for k, (t, e) in ex["track"].items()}
    fails = df[df.tag == "speed"].assign(f=lambda d: d.attempts - d.grasp_ok).groupby(["method", "v"]).f.mean()
    g2 = [1000 * g["err"] for g in ex["trial_m2"]["grasps"]]
    g3 = [1000 * g["err"] for g in ex["trial_m3"]["grasps"]]
    ab_ = ab

    # ============================================================ TITLE PAGE (official sample, only the variable parts change)
    def set_text(par, new):
        runs = par.runs
        runs[0].text = new
        for r in runs[1:]:
            r.text = ""
    paras = d.paragraphs
    set_text(paras[3], TITLE)
    for par in paras[7:10]:
        for r in par.runs:
            r.text = r.text.replace("P00 (Sample)", "P14").replace("G00", "G10").replace("B1", "B4")
    tt = d.tables[0]
    for i, (n, sid) in enumerate(MEMBERS, 1):
        for j, v in enumerate((str(i), n, sid)):
            cell = tt.cell(i, j)
            par = cell.paragraphs[0]
            if par.runs:
                par.runs[0].text = v
                for r in par.runs[1:]:
                    r.text = ""
            else:
                par.add_run(v)

    # ============================================================ FRONT MATTER (roman)
    b.begin_front()
    b.h1("Declaration")
    b.p("We, the undersigned members of Group G10 (Batch B4), declare that the project report entitled \"%s\" is an original piece of work "
        "carried out by us for the course CSAI4014 Algorithms for Intelligent Systems and Robotics. The simulation code, experiments, figures "
        "and text were prepared by the group; every external idea, equation, dataset or figure has been acknowledged and referenced. "
        "The report has not been submitted elsewhere, and no part of it has been copied from another report or from the internet without citation. "
        "We understand that plagiarism leads to cancellation of the project marks as per university rules." % TITLE)
    t = d.add_table(rows=len(MEMBERS) + 1, cols=3)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(("Member", "SAP ID", "Signature")):
        r = t.cell(0, j).paragraphs[0].add_run(h); r.bold = True
    for i, (n, sid) in enumerate(MEMBERS, 1):
        t.cell(i, 0).text = n; t.cell(i, 1).text = sid
    for row in t.rows:
        row.height = Cm(1.1)
        for j, w in enumerate((6.5, 4.0, 5.0)):
            row.cells[j].width = Cm(w)
            for pp in row.cells[j].paragraphs:
                pp.paragraph_format.line_spacing = 1.0; pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par = d.add_paragraph(); par.paragraph_format.space_before = Pt(16); par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.add_run("Date: ____________________\tPlace: Dehradun")

    b.h1("Abstract", True)
    cs_sp = cs_[cs_.tag == "speed"].groupby(["method", "v"]).succ.mean()
    cs_rank = sum(cs_sp.loc[(M3, v_)] >= cs_sp.loc[(M2, v_)] for v_ in (0.10, 0.20, 0.25, 0.30))
    b.p("Quality inspection and sorting on moving conveyors is common in manufacturing, but a robot that must pick moving parts often needs the belt stopped, "
        "which wastes throughput. This project builds a vision-guided sorting arm in CoppeliaSim, the assigned platform, and validates the same software on a second physics engine, MuJoCo. "
        "A vision sensor detects parts by colour and shape: red, green and blue cubes go to colour bins and cylinders, which stand for wrong-shape parts, go to a reject bin. "
        "A SCARA arm with a suction gripper intercepts each part while the belt moves. The Python and OpenCV software contains a colour and shape detector, a Kalman tracker that fuses camera and belt-encoder data, "
        "a D-H based inverse-kinematics intercept planner with minimum-jerk motion, and joint PID control with feedforward. The novelty is a predictive visual-servo descent that keeps correcting the grasp point while the arm moves. "
        "In 560 MuJoCo runs it cut the mean grasp error from %.1f to %.1f mm and kept %.0f %% success at 15 %% belt-speed ripple against %.0f %% for plain prediction; "
        "105 CoppeliaSim runs confirm the ranking in %d of 4 speeds." % (
            sp.loc[M2].err.mean(), sp.loc[M3].err.mean(), rip.loc[(M3, 0.15)].succ, rip.loc[(M2, 0.15)].succ, cs_rank))
    kw = d.add_paragraph()
    kw.add_run("Keywords: ").bold = True
    kw.add_run("vision-guided sorting, quality inspection, SCARA arm, D-H transformations, visual servoing, CoppeliaSim, MuJoCo, OpenCV.")

    b.h1("Index", True)
    p = d.add_paragraph(); p.paragraph_format.line_spacing = 1.0
    add_field(p, 'TOC \\o "1-1" \\h \\z \\u', "Index - update field (right-click, Update Field) if page numbers are missing")

    b.h1("List of Figures", True)
    p = d.add_paragraph(); p.paragraph_format.line_spacing = 1.0
    add_field(p, 'TOC \\h \\z \\c "Figure"', "List of figures - update field")
    b.h1("List of Tables")
    p = d.add_paragraph(); p.paragraph_format.line_spacing = 1.0
    add_field(p, 'TOC \\h \\z \\c "Table"', "List of tables - update field")

    b.h1("Writing Allocation")
    b.p("The writing of this report was divided equally among the five members ({T:alloc}). Each member wrote the sections listed below; every member "
        "also wrote the own reflection in Section 11, filled the own row (with signature) in Section 10, and signed the declaration.")
    alloc = [("Aditya Khatri (500119453)", "Abstract, 1, 2, 12, Appendix"),
             ("Simran Jeet Gill (500122195)", "3, 4, 8"),
             ("Divya Israni (500121919)", "5, 6"),
             ("Amritanshi Singh (500119477)", "7, 9.1-9.2, References"),
             ("Anchit Saxena (500122868)", "9.3-9.8")]
    rows = [["Member", "Name and SAP ID", "Sections written", "Pages"]]
    for i, (n, secs) in enumerate(alloc):
        rows.append(["M%d" % (i + 1), n, secs, pages.get("alloc%d" % i, "")])
    rows.append(["All", "Every member", "10 (own row and signature), 11 (own reflection), Declaration", "-"])
    b.table("alloc", "Writing allocation of the report.", rows, (1.6, 5.0, 6.9, 2.0))

    # ============================================================ MAIN BODY (arabic)
    b.begin_body()

    # ---------------------------------------------------------------- 1
    b.h1("1 Introduction")
    b.h2("1.1 Background")
    b.p("Sorting objects that travel on a conveyor belt is one of the most common applications of industrial robots. Parcel hubs sort "
        "packages by destination, electronics factories separate components by type, food plants grade produce by colour and recycling plants "
        "pick plastics from a moving stream. In all these cases a camera tells the robot what and where the object is, and the robot must "
        "pick it before it leaves the working area. Commercial delta and SCARA robots with line-tracking software reach picking rates of "
        "the order of 100 objects per minute {c:clavel,fanuc,abb}, which is far above what a human operator can do.")
    b.p("Two difficulties make the problem harder than picking a stationary object. First, the object moves between the moment it is "
        "photographed and the moment the gripper touches it. The camera image is processed with a delay of tens of milliseconds, the arm needs "
        "a fraction of a second to move, and the belt speed itself is never perfectly constant. Second, the sensing is noisy: lighting changes, "
        "image noise and calibration errors shift the measured position by millimetres, which is already comparable to the size of a suction cup. "
        "A robot that aims at the position where the object was seen therefore misses it.")
    b.h2("1.2 Why it matters in industry")
    b.p("The simplest industrial remedy is to stop the conveyor while the robot picks, or to place a mechanical fence that holds the object. "
        "This works, but it limits throughput, wears the belt drive and complicates the layout. Continuous line tracking removes these "
        "restrictions and is therefore a key feature of modern pick-and-place cells. Making line tracking accurate and robust against "
        "latency, belt-speed fluctuation and sensor noise decides how fast a cell can run and how many mis-picks occur. "
        "Even a few millimetres of improvement in grasp accuracy allows a smaller suction cup, a higher belt speed or a lower reject rate.")
    b.h2("1.3 Application domain: quality inspection and sorting in manufacturing")
    b.p("The application domain of this project is quality inspection and sorting in manufacturing. On an assembly or packaging line, a camera checks every part that passes: "
        "the colour identifies the variant (for example the colour code of a component or of a cap) and the outline identifies whether the part has the correct shape. "
        "A part with the wrong outline is a reject and must be removed from the flow, while correct parts must be delivered to the bin of their variant. "
        "In our cell the correct parts are cubes in three colours and the rejects are cylinders (25 % of all parts). A robot that does this on a moving belt without stopping it is attractive "
        "because every stop of the belt reduces the output of the whole line. The inspection quality is measured by two error rates: false accepts (a reject delivered to a colour bin) and false rejects (a correct part sent to the reject bin)." )
    b.h2("1.4 Scope of this project")
    b.p("This project (Project P14 of the AISR course) builds a complete vision-guided sorting arm on a conveyor in simulation. The "
        "system recognises parts by colour and shape, tracks them, plans an intercept for a SCARA arm, executes the motion with a joint-level controller and "
        "places each cube into the bin of its colour and each wrong-shape part into the reject bin. Because no physical robot is available, the cell is built in CoppeliaSim {c:coppelia}, "
        "the platform assigned to this project (vision sensor, arm and conveyor, controlled from Python with OpenCV). The same software also runs on a second physics engine, MuJoCo {c:mujoco}, "
        "which is used for the large statistical studies. The report presents the design, the algorithms, "
        "a quantitative comparison of three strategies, stress tests on latency, speed ripple and image noise, and a cross-platform validation. "
        "The remaining sections follow the format of the course: problem statement, literature review, novelty, theory, design, method, "
        "simulation setup, results, contributions, reflections, conclusion and references.")

    # ---------------------------------------------------------------- 2
    b.h1("2 Problem Statement and Objectives")
    b.h2("2.1 Problem definition")
    b.p("A conveyor moves at a nominal speed v between 0.05 and 0.35 m/s. Parts of height 30 mm are placed at random lateral positions and orientations with a random spacing of 0.35 to 0.55 m: "
        "75 % are cubes (side 30 mm) in red, green or blue and 25 % are cylinders (diameter 36 mm) in any of the three colours. A fixed overhead camera observes the belt. "
        "A SCARA arm with a suction gripper stands beside the belt (Figure 1). The task is to detect every part by colour and shape, pick it while the belt keeps moving and to "
        "release cubes into the bin of their colour and cylinders into the reject bin, with the highest possible success rate and throughput. {F:layout} shows the "
        "layout, the camera field of view, the pick window and the four bins.")
    b.fig("layout", "Top view of the sorting cell (belt, camera view, pick window, arm, bins).", 9.0)
    b.h2("2.2 Measurable objectives")
    b.p("The project has five measurable objectives. {T:obj} lists them with the target and the result that verifies each one in Section 9.")
    b.table("obj", "Measurable objectives and the results that verify them.", [
        ["No.", "Objective", "Target", "Achieved (see Section 9)"],
        ["O1", "Detect parts and classify colour and shape in the camera image", ">= 99 % correct at image noise up to 25 gray levels (3 x the default)",
         "colour %.0f %%, shape %.0f %% (noise 25, nominal light)" % (vis[(vis.noise == 25) & (vis.light == 1.0)].acc.iloc[0], vis[(vis.noise == 25) & (vis.light == 1.0)].acc_shape.iloc[0])],
        ["O2", "Grasp accuracy with the moving belt", "mean tip-to-object error <= 5 mm up to 0.35 m/s",
         "%.1f to %.1f mm (M3)" % (sp.loc[M3].err.min(), sp.loc[M3].err.max())],
        ["O3", "Sorting success and robustness", ">= 95 % at v <= 0.20 m/s, and within 10 points of that when belt ripple rises from 3 % to 15 %",
         "%.1f %% at 0.20 m/s; %.1f %% (15 %% ripple) vs %.1f %% (3 %%) (M3)" % (sp.loc[(M3, 0.2)].succ, rip.loc[(M3, 0.15)].succ, rip.loc[(M3, 0.03)].succ)],
        ["O4", "Throughput gain over stop-and-go", ">= 2 x at 0.25 m/s", "%.1f x (%.1f vs %.1f obj/min)" % (
            sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr, sp.loc[(M3, 0.25)].thr, sp.loc[(M1, 0.25)].thr)],
        ["O5", "Inspection quality: wrong accepts and rejects", "false accepts <= 0.1 % and false rejects <= 0.1 % at image noise up to 40",
         "false accepts %d of %d cylinders (%.2f %%), false rejects %d of %d cubes" % (a_fa, a_def, 100.0 * a_fa / a_def, a_fr, a_good)],
    ], (1.1, 5.2, 5.0, 4.2), size=9.5)
    b.h2("2.3 Assumptions and limits")
    b.bullets([
        "The belt is planar and moves along the x-axis; the parts ride the belt without slipping (the conveyor is kinematic: the parts are forced to the belt speed). "
        "The belt speed carries a sinusoidal ripple (default 3 %, tested up to 15 %), and the encoder reading has 3 mm/s noise.",
        "The camera is an overhead orthographic vision sensor (no lens distortion) rendered from the 3-D scene at 1.5 mm/pixel, 20 frames per second and a total latency "
        "of 80 ms. The arm can hide parts from the camera; the carried part is masked in the image.",
        "The arm is a rigid-body model with link masses, inertias and torque limits (45 N m and 18 N m); the vertical axis is ideal in CoppeliaSim and position-controlled in MuJoCo. "
        "The suction gripper is modelled as an attach constraint (CoppeliaSim: the part becomes a child of the tool tip; MuJoCo: a weld) that is switched on when the tip is within 10 mm of the part centre and off at the bin.",
        "Parts are free rigid bodies with contact and friction; a part counts as sorted only if it physically comes to rest inside the correct bin (cubes: colour bin, cylinders: reject bin). Cylinders stand for defective, wrong-shape parts and appear with a probability of 25 %.",
        "Only three colours and two shapes are used and objects never touch each other on the belt. Real-robot issues such as cable slack, vibration, vacuum dynamics "
        "and collision avoidance with other equipment are listed as future work.",
    ])

    # ---------------------------------------------------------------- 3
    b.h1("3 Literature and Industry Review")
    b.p("This section reviews the research and the industrial practice that the project builds on, and identifies the gap that the "
        "project addresses.")
    b.h2("3.1 Robot kinematics, trajectories and control")
    b.p("The kinematics of serial manipulators, including Denavit-Hartenberg parameters, inverse kinematics and the Jacobian, are standard "
        "textbook material {c:sic,craig}. SCARA robots are a planar two-link arm with a vertical axis and are widely used for fast pick-and-place, "
        "because their closed-form inverse kinematics and high horizontal stiffness allow very short cycle times. Minimum-jerk trajectories, "
        "proposed to describe human arm movements by Flash and Hogan {c:flash}, give smooth quintic motion with zero start and end velocity and "
        "acceleration and are common in robot motion planning. Joint-level PID control with model feedforward is still the dominant "
        "industrial servo structure {c:astrom}.")
    b.h2("3.2 Vision and visual servoing")
    b.p("Colour segmentation in HSV space, connected-component analysis and calibration with a planar homography are classical and efficient tools "
        "available in OpenCV {c:cv,hz}. Visual servoing uses image measurements directly inside the control loop. Chaumette and Hutchinson "
        "{c:chau1,chau2} describe position-based and image-based schemes and show that feedback from vision reduces the effect of calibration and model "
        "errors. Corke {c:corke} gives a practical treatment that combines vision and robot control, including moving targets.")
    b.h2("3.3 Estimation and tracking")
    b.p("The Kalman filter {c:kal} is the standard method to estimate position and velocity of a moving target from noisy measurements. "
        "Bar-Shalom et al. {c:bs} treat constant-velocity tracking, data association and delayed measurements in detail. "
        "In conveyor applications the belt encoder gives an almost noise-free speed that can be fused with the camera.")
    b.h2("3.4 Industrial products and standards")
    b.p("Industrial delta robots, introduced by Clavel {c:clavel}, and SCARA robots are combined with line-tracking software in commercial "
        "products such as FANUC iRVision {c:fanuc} and ABB PickMaster {c:abb}. These systems read the belt encoder, trigger the camera at a fixed "
        "belt position and compensate the travelled distance. Performance of such robots is specified with standardised criteria such as "
        "positioning accuracy and repeatability in ISO 9283 {c:iso}. {T:lit} compares these approaches with the approach of this project.")
    b.table("lit", "Comparison of existing approaches with the approach of this project.", [
        ["Approach", "Source", "Strength", "Gap addressed here"],
        ["Stop-and-go conveyor", "common practice", "very accurate and simple", "low throughput (about half in our tests)"],
        ["Encoder-based line tracking", "{c:fanuc,abb}", "high rates, industrial maturity", "closed commercial systems; open comparison with ablations missing"],
        ["Visual servoing theory", "{c:chau1,chau2,corke}", "feedback reduces model error", "rarely combined with intercept planning of a moving, delayed target in an open simulator"],
        ["Kalman target tracking", "{c:kal,bs}", "optimal for linear Gaussian models", "benefit for grasping is rarely quantified against a line fit"],
    ], (3.3, 2.6, 4.2, 5.4), size=9.5)
    b.h2("3.5 Identified gap")
    b.p("Most open educational examples either stop the conveyor or aim at the last measured position, and commercial line-tracking software is closed. "
        "What is missing is a transparent, reproducible study that (i) combines a latency-compensated Kalman tracker, an intercept planner and a "
        "visual-servo descent in one pipeline, (ii) compares it with the stop-and-go and plain predictive baselines over belt speed, latency, "
        "speed ripple and image noise, and (iii) uses an ablation to show which component gives which benefit. This project provides that study.")

    # ---------------------------------------------------------------- 4
    b.h1("4 Novelty")
    b.p("The work is not a new robot but a new combination and analysis that goes beyond the textbook pick-and-place example. "
        "Each point below is supported by measurements from Section 9; {T:nov} summarises the evidence.")
    b.bullets([
        "**N1. Predictive visual-servo descent.** During the 0.2 s descent the tip is driven to the refined grasp point p&#770;(t<sub>g</sub>) through the inverse Jacobian instead of holding the planned point. "
        "At 0.30 m/s with 6 %% ripple and 160 ms latency it lifted success from %.1f %% (without servo) to %.1f %% and cut the mean grasp error from %.1f mm to %.1f mm." % (
            ab.loc["M3 w/o visual servo"].succ, ab.loc[M3].succ, ab.loc["M3 w/o visual servo"].err, ab.loc[M3].err),
        "**N2. Latency-compensated Kalman tracker with encoder fusion.** The belt encoder speed enters the filter as a pseudo-measurement together with a zero lateral velocity constraint. "
        "Against a least-squares line fit it raised success from %.1f %% to %.1f %% and lowered the grasp error from %.2f mm to %.2f mm in the same ablation (significant in the 30-seed confirmation study of Section 9.6). "
        "Its success stays at %.1f %% for every latency from 0 to 320 ms, while the line-fit method falls from %.1f %% to %.1f %%." % (
            ab.loc["M3 w/o Kalman (LS)"].succ, ab.loc[M3].succ, ab.loc["M3 w/o Kalman (LS)"].err, ab.loc[M3].err,
            lat.loc[(M3, 0.32)].succ, lat.loc[(M2, 0.0)].succ, lat.loc[(M2, 0.32)].succ),
        "**N3. Robustness to belt-speed ripple.** At 15 %% ripple the proposed pipeline still sorted %.1f %% of the objects, against %.1f %% for the plain predictive method." % (
            rip.loc[(M3, 0.15)].succ, rip.loc[(M2, 0.15)].succ),
        "**N4. A quantitative three-way study on an open, physics-based simulator.** A MuJoCo cell with a rendered camera, 80 ms latency, encoder noise, a torque-driven arm and rigid-body cubes "
        "was built, and 560 runs compare stop-and-go, plain predictive and the proposed method. Success is counted only if the cube physically rests in the right bin. "
        "The proposed method reaches %.1f times the throughput of stop-and-go at 0.25 m/s." % (sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr),
        "**N5. An honest negative result.** A shortest-cycle-time scheduling rule among feasible targets (SCT) was implemented and tested; it gave no measurable gain over first-in-first-out (%.1f %% in both cases) because all "
        "objects move with the same speed. We report it as a negative finding instead of claiming it." % ab.loc["M3 w/o SCT scheduling"].succ,
    ])
    b.table("nov", "Novelty points and the evidence that supports them.", [
        ["Point", "Evidence", "Where"],
        ["N1 Visual-servo descent", "grasp error %.1f -> %.1f mm; success %.0f -> %.0f %%" % (ab.loc["M3 w/o visual servo"].err, ab.loc[M3].err, ab.loc["M3 w/o visual servo"].succ, ab.loc[M3].succ), "{T:abl}, {F:ablation}"],
        ["N2 Kalman + encoder", "success %.0f -> %.0f %% (30 seeds, p %s); flat versus latency" % (cab["M3 w/o Kalman (LS)"].succ.mean(), cab[M3].succ.mean(), pfmt(welch(cab[M3].succ, cab["M3 w/o Kalman (LS)"].succ))), "{T:conf}, {F:stress}"],
        ["N3 Ripple robustness", "%.0f %% vs %.0f %% success at 15 %% ripple" % (rip.loc[(M3, 0.15)].succ, rip.loc[(M2, 0.15)].succ), "{F:stress}"],
        ["N4 Three-way physics study", "%.1f x throughput over stop-and-go at 0.25 m/s" % (sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr), "{T:speed}, {F:speed}"],
        ["N5 SCT scheduling", "no gain (%.1f %% in both cases)" % ab.loc["M3 w/o SCT scheduling"].succ, "{T:abl}"],
    ], (4.6, 7.0, 3.9), size=9.5)

    # ---------------------------------------------------------------- 5
    b.h1("5 Syllabus Mapping and Theory")
    b.p("{T:syl} maps the AISR topics of Project P14, as named in the Project List, to the sections of this chapter.")
    b.table("syl", "AISR syllabus topics of Project P14 and where they are used.", [
        ["AISR topic (Project List)", "Used for", "Section"],
        ["Robotic Perception", "Rendered camera, homography calibration, HSV colour segmentation, contour-based shape classification, Kalman tracking", "5.5, 5.6"],
        ["D-H Transformations", "D-H table and homogeneous matrices A1 A2 A3 of the SCARA arm, forward and inverse kinematics, Jacobian", "5.1, 5.2"],
        ["Closed-Loop Control", "Joint PID with feedforward, minimum-jerk reference, visual-servo descent (vision in the loop)", "5.3, 5.4"],
        ["Application Domains", "Quality inspection and sorting in manufacturing: colour and shape inspection, reject bin, throughput", "1.3, 5.7"],
    ], (4.4, 9.5, 1.6), size=9.5)
    b.h2("5.1 Kinematics of the SCARA arm")
    b.p("The arm has two revolute joints (q<sub>1</sub>, q<sub>2</sub>) with link lengths L<sub>1</sub> = 0.35 m and L<sub>2</sub> = 0.30 m and a prismatic vertical axis z. "
        "The Denavit-Hartenberg parameters {c:dhorig,craig} are listed in {T:dh}. With the D-H convention each link is described by the homogeneous transformation A<sub>i</sub> = Rot<sub>z</sub>(&#952;<sub>i</sub>) Trans<sub>z</sub>(d<sub>i</sub>) Trans<sub>x</sub>(a<sub>i</sub>) Rot<sub>x</sub>(&#945;<sub>i</sub>):")
    b.table("dh", "Denavit-Hartenberg parameters of the SCARA arm.", [
        ["Joint i", "a<sub>i</sub> (m)", "&#945;<sub>i</sub>", "d<sub>i</sub> (m)", "&#952;<sub>i</sub>"],
        ["1", "0.35", "0", "0", "q1"], ["2", "0.30", "0", "0", "q2"], ["3", "0", "0", "-z", "0 (prismatic)"]],
        (2.5, 3.0, 3.0, 3.0, 3.5), size=10, align_left=False)
    b.eq("A<sub>i</sub> = [cθ<sub>i</sub>  −sθ<sub>i</sub>cα<sub>i</sub>  sθ<sub>i</sub>sα<sub>i</sub>  a<sub>i</sub>cθ<sub>i</sub>;  sθ<sub>i</sub>  cθ<sub>i</sub>cα<sub>i</sub>  −cθ<sub>i</sub>sα<sub>i</sub>  a<sub>i</sub>sθ<sub>i</sub>;  0  sα<sub>i</sub>  cα<sub>i</sub>  d<sub>i</sub>;  0  0  0  1]")
    b.p("For the SCARA table (all &#945;<sub>i</sub> = 0) the chain T<sub>0</sub><super>3</super> = A<sub>1</sub> A<sub>2</sub> A<sub>3</sub> gives the pose of the tool frame:")
    b.eq("T = [c₁₂  −s₁₂  0  L₁c₁ + L₂c₁₂;  s₁₂  c₁₂  0  L₁s₁ + L₂s₁₂;  0  0  1  −z;  0  0  0  1],     c₁₂ = cos(q₁ + q₂)")
    b.p("Its position column is the forward kinematics of the tip, and its rotation part shows that the tool turns by q<sub>1</sub> + q<sub>2</sub> about the vertical axis. A unit test confirms that the matrix chain equals the closed-form tip position for 100 random poses:")
    b.eq("x = L₁ cos q₁ + L₂ cos(q₁ + q₂),     y = L₁ sin q₁ + L₂ sin(q₁ + q₂)")
    b.p("For a desired tip point (x, y) the closed-form inverse kinematics gives two elbow solutions:")
    b.eq("cos q₂ = (x² + y² − L₁² − L₂²) / (2 L₁ L₂),     q₂ = ± arccos(cos q₂)")
    b.eq("q₁ = atan2(y, x) − atan2(L₂ sin q₂,  L₁ + L₂ cos q₂)")
    b.p("The solution within the joint limits (|q<sub>1</sub>| <= 2.97 rad, |q<sub>2</sub>| <= 2.53 rad) that is closest to the present joint angles is selected. "
        "A target outside the reachable annulus (|c<sub>2</sub>| > 1) is rejected.")
    b.h2("5.2 Jacobian and resolved-rate control")
    b.p("The velocity of the tip is related to the joint rates by the Jacobian J(q), with det J = L<sub>1</sub> L<sub>2</sub> sin q<sub>2</sub>. Away from the singularities "
        "(q<sub>2</sub> = 0, &#960;), a desired tip velocity is converted to joint rates by:")
    b.eq("dq/dt = J(q)⁻¹ v<sub>tip</sub>")
    b.p("This relation drives the visual-servo descent in Section 7.5. In the workspace used here |q<sub>2</sub>| stays between 1.1 and 2.4 rad, so the Jacobian is well conditioned.")
    b.h2("5.3 Minimum-jerk trajectory")
    b.p("Each point-to-point motion from q<sub>0</sub> to q<sub>1</sub> in time T uses the quintic blend {c:flash}:")
    b.eq("q(t) = q<sub>s</sub> + (q<sub>e</sub> − q<sub>s</sub>) s(τ),     s(τ) = 10τ³ − 15τ⁴ + 6τ⁵,     τ = t / T")
    b.p("Its peak speed is 1.875 |&#916;q|/T and its peak acceleration 5.77 |&#916;q|/T<super>2</super>. The shortest feasible duration for the limits &#969;<sub>max</sub> = 6 rad/s and "
        "&#945;<sub>max</sub> = 40 rad/s<super>2</super> is therefore:")
    b.eq("T<sub>min</sub> = max<sub>joints</sub> [ 1.875 |Δq| / ω<sub>max</sub> ,  √(5.77 |Δq| / α<sub>max</sub>) ]")
    b.h2("5.4 Joint dynamics and PID control")
    b.p("In the simulator the arm obeys the rigid-body equation M(q) d<super>2</super>q/dt<super>2</super> + C(q, dq/dt) + B dq/dt = &#964; with torque saturation. The controller treats each joint as J d<super>2</super>q/dt<super>2</super> + B dq/dt = &#964; with the constant nominal inertia J<sub>n</sub> (taken from the mass matrix at a representative pose, %.2f and %.3f kg m<super>2</super>) and damping B<sub>n</sub>, and uses the "
        "PID law with feedforward {c:astrom}, where q<sub>r</sub> is the reference trajectory and e the joint error:" % J_NOM)
    b.eq("τ = J<sub>n</sub> d²q<sub>r</sub>/dt² + B<sub>n</sub> dq<sub>r</sub>/dt + K<sub>p</sub> e + K<sub>d</sub> de/dt + K<sub>i</sub> ∫ e dt")
    b.p("The gains follow from a desired natural frequency &#969;<sub>n</sub> = 50 rad/s and damping &#950; = 1: K<sub>p</sub> = J<sub>n</sub>&#969;<sub>n</sub><super>2</super>, "
        "K<sub>d</sub> = 2&#950;&#969;<sub>n</sub>J<sub>n</sub>, K<sub>i</sub> = 0.2 K<sub>p</sub>&#969;<sub>n</sub>. The inertia seen by a joint changes with the arm pose and the two joints are coupled, so the feedforward is imperfect and the PID has to correct the rest.")
    b.h2("5.5 Camera model, calibration, colour and shape detection")
    b.p("A planar scene seen by a camera is related to the image by a homography H. With four marker points of known world position the matrix H is "
        "estimated by the direct linear transform {c:hz}, which OpenCV implements in getPerspectiveTransform {c:cv}:")
    b.eq("(X, Y, 1)<super>T</super>  ∼  H (u, v, 1)<super>T</super>")
    b.p("Colour is classified in HSV space, because hue is almost independent of brightness: red is the hue range below 8 or above 172, green 45 to 85 and "
        "blue 95 to 125 (OpenCV scale 0 to 179), each with saturation above 90 and value above 50. Morphological opening removes isolated noise pixels, and the centroid "
        "of each connected component gives the part position with sub-pixel accuracy {c:gw}.")
    b.p("The shape is decided from the contour of each component. A circle has the same distance from its centroid to every contour point, a square has "
        "distances between half the side and half the diagonal. The coefficient of variation of the contour radius r is therefore a simple, rotation-invariant shape feature:")
    b.eq("CV = σ<sub>r</sub> / μ<sub>r</sub>,     circle if CV < 0.06,  square otherwise")
    b.p("At 1.5 mm/pixel the measured values are about 0.02 for the cylinders (pixels spread over 0.019 to 0.026) and 0.09 to 0.11 for the cubes, so the threshold lies far from both classes. "
        "At 2.5 mm/pixel the two classes overlapped, which is why the camera resolution was raised.")
    b.h2("5.6 Kalman filter for moving objects")
    b.p("The state of each object is x = [x, y, v<sub>x</sub>, v<sub>y</sub>]<super>T</super> with a constant-velocity model and white acceleration noise (&#963;<sub>a</sub> = 0.08 m/s<super>2</super>) {c:kal,bs}:")
    b.eq("x<sub>k</sub> = F x<sub>k−1</sub> + w,     F = [1 0 Δt 0; 0 1 0 Δt; 0 0 1 0; 0 0 0 1]")
    b.eq("K = P (P + R)⁻¹,     x ← x + K (z − x),     P ← (I − K) P")
    b.p("The measurement z contains the camera position (&#963; = 3 mm), the encoder belt speed (&#963; = 4 mm/s) and a zero lateral velocity (&#963; = 2 mm/s), so that H = I. "
        "A measurement taken at time t<sub>f</sub> reaches the filter at t<sub>f</sub> + 80 ms. The filter state therefore lags and is extrapolated to the present time: "
        "p(t) = p&#770; + v&#770; (t - t<sub>f</sub>). This is the latency compensation.")
    b.h2("5.7 Intercept planning")
    b.p("For a tracked object the planner searches the earliest arrival time t<sub>a</sub> (in steps of 20 ms) for which the grasp point p(t<sub>a</sub> + T<sub>d</sub>) lies in the pick window "
        "(-0.30 m <= x <= 0.30 m), the inverse kinematics has a solution q<sub>g</sub> and the minimum-jerk motion from the present pose fits in the available time:")
    b.eq("t<sub>a</sub> − t<sub>plan</sub> ≥ T<sub>min</sub>(q<sub>now</sub>, q<sub>g</sub>) + 10 ms")
    b.p("Here T<sub>d</sub> = 0.2 s is the descent time. If no such t<sub>a</sub> exists before the object leaves the window, the object is skipped.")

    # ---------------------------------------------------------------- 6
    b.h1("6 System Design")
    b.h2("6.1 Architecture")
    b.p("{F:block} shows the architecture. The simulated belt carries the cubes and provides an encoder reading. The camera renders an image 20 times per second; the vision "
        "module extracts blobs, which arrive at the tracker 80 ms later. The tracker maintains one Kalman filter per part. When the arm is idle, "
        "the planner selects a target and an intercept time, the trajectory generator produces joint references and the PID controllers drive the joints. "
        "During the descent the visual-servo block corrects the references with the latest estimate; in CoppeliaSim the joint controller runs in joint-callback scripts inside the physics step. After the grasp the arm carries the part to the bin of its recognised colour (cubes) or to the reject bin (cylinders) and releases it.")
    b.fig("block", "Block diagram of the complete system.", 10.0)
    b.h2("6.2 Robot and environment model")
    b.fig("scene", "The simulated cell: (a) CoppeliaSim (assigned platform) and (b) MuJoCo; SCARA arm, conveyor, three colour bins and the dark reject bin.", 9.5)
    b.p("{F:scene} shows the simulated cell. The arm is mounted at the origin of the world frame beside the belt. The belt occupies 0.26 m <= y <= 0.40 m and moves in the +x direction. The camera field of view "
        "covers -0.85 m <= x <= 0.35 m and 0.20 m <= y <= 0.46 m, which gives an image of 800 x 173 pixels at 1.5 mm per pixel. Three bins for red, green and blue cubes lie on the opposite "
        "side of the arm at (-0.25, -0.30), (0, -0.35) and (0.25, -0.30) m; the reject bin for cylinders is at (0.48, -0.12) m. The principal parameters are given in {T:par}.")
    b.table("par", "Principal parameters of the simulated cell.", [
        ["Group", "Parameter", "Value"],
        ["Arm", "Link lengths L1, L2", "0.35 m, 0.30 m"],
        ["Arm", "Joint limits q1, q2", "+/- 2.97 rad, +/- 2.53 rad"],
        ["Arm", "Speed / acceleration limit (planner)", "6 rad/s, 40 rad/s^2"],
        ["Arm", "Torque limits", "45 N m, 18 N m"],
        ["Arm", "Link masses (link 1 / link 2 / tool)", "1.8 kg / 1.1 kg / 0.17 kg"],
        ["Arm", "Nominal inertia used by controller", "%.2f, %.3f kg m^2" % J_NOM],
        ["Physics", "Engine, time step", "CoppeliaSim 4.10.0 (Bullet), 5 ms; MuJoCo %s, 2 ms" % mujoco.__version__],
        ["Gripper", "Suction capture radius", "10 mm"],
        ["Conveyor", "Speed range, ripple, encoder noise", "0.05-0.35 m/s, 3 % (0-15 %), 3 mm/s"],
        ["Objects", "Cube side, cylinder diameter, share of cylinders", "30 mm, 36 mm, 25 %"],
        ["Objects", "Spacing on the belt", "0.35-0.55 m"],
        ["Camera", "Resolution, frame rate, latency", "1.5 mm/pixel (800 x 173), 20 fps, 80 ms"],
        ["Camera", "Image noise, calibration noise", "8 gray levels, 0.8 mm"],
        ["Camera", "Type", "orthographic vision sensor, rendered from the 3-D scene"],
        ["Control", "Joint loop bandwidth, damping", "50 rad/s, 1 (both platforms)"],
    ], (2.6, 7.2, 5.7), size=9.5)
    b.h2("6.3 Sensors and actuators")
    b.p("The sensors are the overhead camera (position, colour and shape), the belt encoder (speed) and the joint encoders of the arm (angles and rates, used by the controller and for "
        "forward kinematics). The actuators are the two joint motors, the vertical axis and the suction gripper. The camera and encoder deliver data at 20 Hz and the arm controller runs at 500 Hz in MuJoCo (2 ms) and at 200 Hz in CoppeliaSim (5 ms).")
    b.h2("6.4 Software architecture")
    b.p("The code is organised in small modules, each of which belongs to one team member (Section 10): arm.py (kinematics, D-H matrices, trajectory, PID law), vision.py (detector, shape classifier, homography), "
        "tracker.py (Kalman and line-fit tracking), sim.py (planner and arm state machine), cs_world.py, simcs.py and cs_launch.py (CoppeliaSim scene, closed loop, launcher), mj_world.py and sim3d.py (MuJoCo), "
        "the experiment, demo and render scripts, and tests/. The arm state machine has six states: IDLE, MOVE, DESCEND, TRANSPORT, RELEASE and RISE (after a miss).")

    # ---------------------------------------------------------------- 7
    b.h1("7 Methodology and Algorithms")
    b.h2("7.1 Overall procedure")
    b.p("The complete cycle is summarised in the flowchart of {F:flow}. Every 50 ms a frame is rendered and processed; every 20 ms, if the arm is idle, the planner is invoked. Everything else runs at 2 ms.")
    b.fig("flow", "Flowchart of the proposed method (M3), from image capture to bin release.", 7.2)
    b.h2("7.2 Vision pipeline")
    b.p("The camera image is rendered by an orthographic vision sensor above the belt (CoppeliaSim; MuJoCo uses an equivalent camera), showing parts, belt and the arm itself, which can occlude parts. The brightness is scaled by a time-varying lighting factor "
        "(+/- 12 %) and Gaussian pixel noise (default 8 gray levels) is added. The detector blurs the image, converts it to HSV, applies the three colour "
        "thresholds, removes noise by opening, labels connected components and discards blobs that are too small or touch the image border (partly visible parts would bias the centroid). "
        "For every remaining blob the contour is extracted and the shape feature CV of Section 5.5 gives the class square or circle. "
        "{F:camera} shows an example frame and the detections. Pixel centroids are mapped to world coordinates with the homography of Section 5.5.")
    b.fig("camera", "(a) Camera frame rendered by MuJoCo and (b) detected blobs with colour, shape and centroid.", 9.0)
    b.h2("7.3 Tracking and prediction")
    b.p("Detections are assigned to existing tracks by nearest neighbour within a 50 mm gate; unassigned detections start new tracks and tracks that are not updated for 0.6 s are deleted. "
        "A track is confirmed after three detections and the class label is decided by majority vote. Two estimators were implemented. The baseline fits a straight line to the last twelve measurements "
        "by least squares. The proposed estimator is the Kalman filter of Section 5.6 with encoder fusion and latency compensation.")
    b.h2("7.4 Planning algorithm")
    b.p("Algorithm 1 gives the planner in pseudo-code.")
    b.code(["Algorithm 1: Intercept planning (called when the arm is idle)",
            "1  for each confirmed, unclaimed track tr with fails < 2:",
            "2      (x, y, vx, vy) <- tr.estimate(t_now)      # latency compensated",
            "3      if x outside [-0.30, 0.30]: continue",
            "4      for t_a = t_now + 0.10 ... in steps of 0.02 s:",
            "5          g <- (x, y) + (vx, vy) * (t_a + T_d - t_now)   # grasp point",
            "6          if g.x > 0.30: break                  # object would be lost",
            "7          q_g <- IK(g) nearest to q_now; if none: continue",
            "8          if t_a - t_now - t_plan >= T_min(q_now, q_g) + 0.01:",
            "9              add (tr, t_a, g, q_g) to candidates; break",
            "10 pick the candidate with the shortest cycle time (FIFO: the first)",
            "11 claim it; plan minimum-jerk motion q_now -> q_g, arrival at t_a"])
    b.h2("7.5 Predictive visual-servo descent")
    b.p("After arrival at the hover height the tip descends 0.10 m in T<sub>d</sub> = 0.2 s. In the plain predictive method the arm holds the planned point during the descent. "
        "In the proposed method the references are updated every 2 ms. The latest Kalman estimate is extrapolated to the grasp time t<sub>g</sub> and the tip is driven towards it with a proportional "
        "law in Cartesian space and the inverse Jacobian (Algorithm 2, K<sub>t</sub> = 15 s<super>-1</super>). The point p&#770;(t<sub>g</sub>) is used rather than the current object position, because the arm arrives ahead of the cube "
        "and waits for it, as shown in the traces of Section 9.")
    b.code(["Algorithm 2: Visual-servo descent (every 2 ms until t_g = t_arrival + T_d)",
            "1  z <- minimum-jerk descent from z_hover to z_pick",
            "2  p_goal <- tr.estimate(t_g)        # predicted cube position at contact",
            "3  p_tip  <- FK(q_measured)",
            "4  v_cmd  <- Kt * (p_goal - p_tip)",
            "5  dq     <- J(q_ref)^-1 * v_cmd",
            "6  q_ref  <- q_ref + dq * dt ;  qdot_ref <- dq",
            "7  at t_g: grasp succeeds if |p_tip - p_cube| < r_cap = 10 mm"])
    b.h2("7.6 Controller design and tuning")
    b.p("The gains were chosen from the pole-placement formulas of Section 5.4 with &#969;<sub>n</sub> = 50 rad/s (about eight times faster than the minimum-jerk motions). The gains were verified with a 0.5 rad step "
        "and a 1.2 rad minimum-jerk tracking test (Section 9.1). With a 2 ms step and semi-implicit Euler integration the loop is stable (&#969;<sub>n</sub> dt = 0.1). The integral term is clamped to avoid wind-up "
        "when the torque saturates.")
    b.h2("7.7 Methods compared")
    b.p("{T:meth} lists the three strategies and the three variants that are used in the ablation study.")
    b.table("meth", "The three strategies and the ablated variants.", [
        ["Method", "Belt", "Estimator", "Descent", "Scheduling"],
        ["M1 Stop-and-go", "stops while an object is in the window", "mean of still-belt detections", "hold", "nearest first"],
        ["M2 Predictive", "moves", "least-squares line fit", "hold planned point", "FIFO"],
        ["M3 Proposed", "moves", "Kalman + encoder", "visual servo", "shortest cycle (SCT)"],
        ["M3 w/o visual servo", "moves", "Kalman + encoder", "hold", "SCT"],
        ["M3 w/o Kalman", "moves", "line fit", "visual servo", "SCT"],
        ["M3 w/o SCT", "moves", "Kalman + encoder", "visual servo", "FIFO"],
    ], (3.6, 3.6, 3.4, 2.6, 2.3), size=9.5)

    # ---------------------------------------------------------------- 8
    b.h1("8 Simulation Setup")
    b.h2("8.1 Platform")
    b.p("The assigned platform of Group G10 is CoppeliaSim {c:coppelia} (vision sensor, arm and conveyor) with Python and OpenCV. The cell is built in CoppeliaSim V4.10.0 (EDU) with the Bullet physics engine "
        "and is controlled from Python through the ZMQ remote API. The software of the project (detector, tracker, planner, visual servo) is plain Python 3 with NumPy, OpenCV {c:cv} and SciPy and is identical on both platforms. "
        "CoppeliaSim hosts the live demonstration, the video and a cross-platform validation study. One CoppeliaSim trial of 12 parts takes about 40 to 80 s of computing time, about five times longer than in MuJoCo {c:mujoco}, "
        "so the large statistical studies were run in MuJoCo. Running the identical software on two independent physics engines also shows that the results are not an artefact of one engine (Section 9.7).")
    b.p("In CoppeliaSim the scene is created from Python with the remote API (cs_world.py): static shapes for the belt and the four bins, two dynamic revolute joints with their links, a kinematic vertical axis with the suction tool, "
        "an orthographic vision sensor, and parts created at the start of the belt. The joint-level PID with feedforward of Section 5.4 is evaluated in a joint-callback script inside every physics step (exact timing), "
        "and one batched call per step sets the references and moves the conveyor parts. In MuJoCo the arm, the belt and the bins are defined in an MJCF model (mj_world.py) and the torque controller runs in Python. "
        "Randomness (part colour, shape, position, orientation, spacing, image noise and calibration error) is controlled by an integer seed so that the same seeds are used for all methods. {T:map} maps the elements named in the Project List to the implementation.")
    b.table("map", "Elements of the CoppeliaSim set-up named in the Project List and their implementation.", [
        ["Element (Project List)", "Implementation in this project"],
        ["Vision sensor", "Orthographic vision sensor, 800 x 173 px, 20 fps; the image goes to OpenCV, noise and latency are added in Python"],
        ["Arm", "SCARA: base, two dynamic revolute joints, links, kinematic vertical axis, suction tool; D-H table of {T:dh}; PID + feedforward in a joint callback"],
        ["Conveyor", "Belt shape; the parts are moved along it every step with speed ripple (kinematic); the encoder value goes to the tracker"],
        ["Parts and bins", "Cubes and cylinders as Bullet rigid bodies; four open bins; success read from the physics state"],
        ["Python + OpenCV", "Detector, tracker, planner and visual servo: the same code on CoppeliaSim and MuJoCo"],
    ], (4.0, 11.5), size=9.5)
    b.h2("8.2 Experiments")
    b.bullets([
        "**MuJoCo belt speed study:** 3 methods x 7 speeds (0.05 to 0.35 m/s) x 10 seeds, 15 parts per run (150 parts per point).",
        "**MuJoCo stress tests at 0.25 m/s (M2 and M3):** latency 0 to 320 ms, speed ripple 0 to 15 %, image noise 4 to 60 gray levels; ablation at 0.30 m/s with 6 % ripple and 160 ms latency (five variants).",
        "**MuJoCo confirmation study:** 30 new seeds for the key comparisons (330 runs, Section 9.6).",
        "**CoppeliaSim cross-platform study:** 105 runs with 5 seeds and 12 parts per run on 8 headless CoppeliaSim instances: 3 methods at 4 speeds, stress tests at 15 % ripple and 320 ms latency, and the ablation (Section 9.7).",
        "**Controller and vision studies:** joint step response and minimum-jerk tracking on both platforms, and a detection study with %d parts per condition." % vis.objects.iloc[0],
    ])
    b.h2("8.3 Performance measures")
    b.p("A part counts as sorted correctly only if, after the simulation, it physically rests inside its correct bin (cubes: the bin of their true colour; cylinders: the reject bin) (the position is read from the physics engine, not from the planner). The success rate is the percentage of all parts that were sorted correctly; "
        "parts that left the pick window are counted as lost. The throughput is the number of correctly sorted parts per minute of simulated time, "
        "measured until the last part is resolved. The grasp error is the tip-to-part distance at the moment of suction.")
    b.h2("8.4 How to run the code")
    b.p("{T:sim} lists the commands that reproduce every result of this report.")
    b.table("sim", "Files of the repository and commands to reproduce all results.", [
        ["Command", "Purpose"],
        ["pip install -r requirements.txt", "install the libraries (CoppeliaSim: docs/COPPELIASIM_SETUP.md)"],
        ["python demo_cs.py [--method m2 --speed 0.3]", "live demo in the CoppeliaSim window (used in the video)"],
        ["python render_cs.py --stills", "CoppeliaSim stills and MP4 video"],
        ["python demo.py  /  python render_3d.py", "the same demo and renders in MuJoCo"],
        ["python tests/test_core.py", "11 unit tests"],
        ["python experiments.py [--confirm]", "560 + 330 MuJoCo runs -> results/*.csv (about 25 min)"],
        ["python experiments_cs.py", "105 CoppeliaSim runs -> results/cs_runs.csv (about 15 min)"],
        ["python extra_experiments.py  /  python extra_cs.py", "controller and vision studies on both platforms"],
        ["python make_figures.py, python report/build_report.py", "regenerate figures, rebuild the report"],
    ], (6.6, 8.9), size=9.5)

    b.h2("8.5 Model validation and tests")
    b.p("Eleven unit tests (tests/test_core.py) check the building blocks, and the checks of {T:val} show that the planner, the D-H model and both physics models use the same geometry "
        "and that the detector, the tracker and the minimum-jerk planner behave as specified.")
    b.table("val", "Validation checks of the simulation.", [
        ["Check", "Result"],
        ["IK followed by FK on 500 random targets", "max error < 1e-9 m"],
        ["Analytic vs finite-difference Jacobian", "max difference < 1e-5"],
        ["MuJoCo tool position vs analytic FK (50 poses)", "max error < 1e-9 m"],
        ["Minimum-jerk peak speed factor", "1.875 (exact)"],
        ["Kalman velocity estimate, 0.2 m/s belt", "error < 0.01 m/s"],
        ["D-H chain A1 A2 A3 vs closed-form tip position (100 poses)", "max error < 1e-12 m"],
        ["CoppeliaSim tool tip vs analytic FK (q1 = 0.8, q2 = 1.0)", "(0.176, 0.543) m in both"],
        ["CoppeliaSim vision sensor vs true part positions (4 parts)", "centroid error <= 1.5 mm, colour and shape all correct"],
        ["Detector on a rendered scene (squares, circles, 3 colours)", "colour and shape all correct"],
        ["Parts counted as sorted", "verified from physics state, not from planner log"],
    ], (9.0, 6.5), size=9.5)

    # ---------------------------------------------------------------- 9
    b.h1("9 Results and Analysis")
    b.h2("9.1 Joint controller")
    b.p("{F:control}(a) shows the response of joint 1 of the MuJoCo arm to a 0.5 rad step. A proportional controller alone oscillates (%.0f %% overshoot, not settled after 0.6 s). "
        "PD control has %.1f %% overshoot and settles within 2 %% in %.2f s; PID adds an integral term with %.1f %% overshoot and %.2f s settling, which removes the steady-state error that joint damping and coupling torques would otherwise leave while the arm holds the hover pose. "
        "For the minimum-jerk motions used in the cell, panel (b), the maximum tracking error is %.1f mrad with PD, %.1f mrad with PID and %.1f mrad with PID plus feedforward. "
        "The feedforward cuts the PID error by %.0f %%; its peak error equals that of PD, because the nominal inertia is only an approximation of the pose-dependent inertia of the coupled arm. "
        "At the 0.4 m reach, %.1f mrad corresponds to %.1f mm at the tip." % (
            step["P"][0], max(step["PD"][0], 0.0), step["PD"][1], step["PID"][0], step["PID"][1], trk["PD"], trk["PID"], trk["PID+FF"],
            100 * (1 - trk["PID+FF"] / trk["PID"]), trk["PID+FF"], 0.4 * trk["PID+FF"]))
    b.fig("control", "Joint 1 controller: (a) step response; (b) tracking error (solid: MuJoCo, dashed: CoppeliaSim).", 10.8)
    if cs_ctl is not None:
        def _ov(y_):
            return 100 * (max(y_) - 0.5) / 0.5
        trk_cs = {k_: 1000 * float(abs(e_).max()) for k_, (t_, e_) in cs_ctl["track"].items()}
        b.p("The same tests in CoppeliaSim (dashed curves; the PID + feedforward law runs in a joint-callback script every 5 ms) give the same qualitative result: the P-only loop oscillates (%.0f %% overshoot), PD has %.1f %% overshoot and PID %.1f %%. "
            "The peak tracking error of the minimum-jerk motion is %.1f mrad with PD, %.1f mrad with PID and %.1f mrad with PID plus feedforward, about three times larger than in MuJoCo because the loop is sampled at 200 Hz instead of 500 Hz and the Bullet joints are softer. "
            "This corresponds to %.1f mm at the 0.4 m reach and explains the larger grasp error of the CoppeliaSim cell." % (
                _ov(cs_ctl["step"]["P"][1]), max(_ov(cs_ctl["step"]["PD"][1]), 0.0), _ov(cs_ctl["step"]["PID"][1]), trk_cs["PD"], trk_cs["PID"], trk_cs["PID+FF"], 0.4 * trk_cs["PID+FF"]))
    b.h2("9.2 Vision module")
    v40 = vis[(vis.noise == 40) & (vis.light == 1.0)]
    v60 = vis[(vis.noise == 60) & (vis.light == 0.7)]
    v80 = vis[(vis.noise == 80) & (vis.light == 0.7)]
    b.p("The detector was tested on %d parts per condition for image noise from 4 to 80 gray levels and for three lighting levels ({F:vision}). Colour classification is the more robust feature: it is "
        "%.0f %% correct at every noise level up to 60 and %.0f %% at noise 80 with light x0.7. Shape classification is the limiting step: it is %.0f %% correct at noise 25 and %.1f %% at noise 40 with nominal light, "
        "and falls to %.0f %% at noise 60 and %.0f %% at noise 60 with light x0.7, because noise makes the contour of a cylinder ragged and raises its radius variation CV towards the threshold. "
        "A 5 x 5 Gaussian pre-blur limits this effect; the default noise (8 gray levels) is far below these limits. "
        "The position RMSE is about %.1f mm. Without calibration noise it is 1.4 mm (the detector itself: 1.5 mm pixels, shadows and the part edges); the 0.8 mm noise on the four calibration markers raises it to 2.2 mm, so the calibration is the larger part of the perception error; image noise adds almost nothing. Missed detections are bridged by the tracker, which also decides colour and shape by majority vote over all frames." % (
            vis.objects.iloc[0], vis[(vis.noise == 60) & (vis.light == 1.0)].acc.iloc[0], v80.acc.iloc[0],
            vis[(vis.noise == 25) & (vis.light == 1.0)].acc_shape.iloc[0], v40.acc_shape.iloc[0], vis[(vis.noise == 60) & (vis.light == 1.0)].acc_shape.iloc[0], v60.acc_shape.iloc[0],
            vis[vis.noise == 8].rmse_mm.mean()))
    b.fig("vision", "Vision accuracy versus image noise for three lighting levels.", 9.5)
    b.h2("9.3 Comparison of the three strategies")
    rows = [["Speed (m/s)", "M1 succ. %", "M1 obj/min", "M2 succ. %", "M2 err mm", "M3 succ. %", "M3 obj/min", "M3 err mm"]]
    for v in (0.05, 0.10, 0.15, 0.20, 0.25, 0.30, 0.35):
        rows.append(["%.2f" % v, "%.1f" % sp.loc[(M1, v)].succ, "%.1f" % sp.loc[(M1, v)].thr, "%.1f" % sp.loc[(M2, v)].succ, "%.1f" % sp.loc[(M2, v)].err,
                     "%.1f" % sp.loc[(M3, v)].succ, "%.1f" % sp.loc[(M3, v)].thr, "%.1f" % sp.loc[(M3, v)].err])
    b.p("{T:speed} and {F:speed} compare the strategies over belt speed (mean of 10 seeds, error bars: one standard deviation). "
        "The stop-and-go cell (M1) sorts 100 %% of the cubes at every speed, but its throughput is only %.1f to %.1f objects per minute, because the whole belt "
        "waits for each pick. The moving-belt methods are much faster: at 0.25 m/s the proposed method reaches %.1f objects per minute, %.1f times the stop-and-go rate (objective O4). "
        "Up to 0.10 m/s M2 and M3 sort at least 99 %% of the cubes. From 0.15 m/s on M2 starts to fail (%.1f %% at 0.15 m/s, %.1f %% at 0.20 m/s) because its aim point is off by 4 to 5 mm and the suction cup captures only 10 mm, "
        "so it makes %.1f failed grasp attempts per run at 0.25 m/s, while M3 makes %.1f. M3 sorts %.1f %% at 0.20 m/s (objective O3). "
        "The mean grasp error of M2 is %.1f to %.1f mm, that of M3 %.1f to %.1f mm (objective O2). Even the stop-and-go cell has %.1f mm; this is the perception floor (calibration and centroid error, Section 9.2), so no method can be much better. "
        "Above 0.25 m/s the success of both moving-belt methods drops (M3: %.1f %% at 0.30 m/s and %.1f %% at 0.35 m/s) for a different reason than accuracy: the arm needs about 2 s for a full pick-and-place cycle, "
        "so objects arrive faster than they can be served and some leave the pick window. This is a capacity limit of the single arm; M3 still makes no failed grasps." % (
            sp.loc[M1].thr.min(), sp.loc[M1].thr.max(), sp.loc[(M3, 0.25)].thr, sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr,
            sp.loc[(M2, 0.15)].succ, sp.loc[(M2, 0.20)].succ, fails.loc[(M2, 0.25)], fails.loc[(M3, 0.25)], sp.loc[(M3, 0.2)].succ,
            sp.loc[M2].err.min(), sp.loc[M2].err.max(), sp.loc[M3].err.min(), sp.loc[M3].err.max(), sp.loc[M1].err.mean(),
            sp.loc[(M3, 0.30)].succ, sp.loc[(M3, 0.35)].succ))
    b.table("speed", "Results versus belt speed (mean of 10 runs with 15 objects each).", rows, (1.9, 1.9, 2.0, 1.9, 1.9, 1.9, 2.0, 2.0), size=9, align_left=False)
    b.fig("speed", "(a) Sorting success, (b) throughput and (c) grasp error versus belt speed.", 11.0)
    b.p("**Inspection quality.** In all runs with image noise up to 40 (the 560-run study and the 330-run confirmation study together: %d cylinders and %d cubes) %d defective part was accepted into a colour bin (false accept, %.2f %%) and "
        "%d good cubes were sent to the reject bin; no cube was delivered to a wrong colour bin (objective O5). The single false accept happened with the plain predictive method M2 at 15 %% belt-speed ripple: "
        "the track followed a blue cube but the suction cup reached a neighbouring blue cylinder, so the bin was chosen from the wrong track (verified by replaying the run). All other sorting failures are parts that were not handled in time (lost from the belt), never a wrong classification. "
        "Only in the extreme-noise runs (noise 60, %d parts, %d cylinders) did the shape classifier fail: %d cylinders were accepted into a colour bin. Colour classification was correct in all runs." % (
            a_def, a_good, a_fa, 100.0 * a_fa / a_def, a_fr, h_n, h_def, h_fa))
    b.h2("9.4 Robustness to latency, speed ripple and image noise")
    nz2 = [noi.loc[(M2, k)].succ for k in (4.0, 8.0, 20.0, 40.0, 60.0)]
    b.p("{F:stress} shows the stress tests at 0.25 m/s. When the vision latency is increased from 0 to 320 ms, the success of M2 falls from %.1f %% to %.1f %% and its grasp error rises from %.1f to %.1f mm; "
        "M3 keeps its success at %.1f %% at every latency and its error between %.1f and %.1f mm, because the latency is compensated in the filter and the descent is corrected with the latest estimate. The speed ripple has the strongest effect: "
        "at 15 %% ripple M2 sorts only %.1f %% of the cubes, with a grasp error of %.1f mm, while M3 still sorts %.1f %% with %.1f mm (objective O3: %.1f points lower than at 3 %% ripple). "
        "A speed error of 1 %% over a 1 s look-ahead at 0.25 m/s shifts the target by 2.5 mm, so a ripple of this size is critical for a 10 mm suction cup. "
        "Image noise up to 40 gray levels changes neither method beyond the run-to-run scatter (M3: %.1f %% at every level; M2: %.1f to %.1f %%); at noise 60 the success falls to %.1f %% (M3) and %.1f %% (M2) because of the shape errors described in Section 9.2." % (
            lat.loc[(M2, 0.0)].succ, lat.loc[(M2, 0.32)].succ, lat.loc[(M2, 0.0)].err, lat.loc[(M2, 0.32)].err, lat.loc[(M3, 0.32)].succ, lat.loc[(M3, 0.32)].err, lat.loc[(M3, 0.0)].err,
            rip.loc[(M2, 0.15)].succ, rip.loc[(M2, 0.15)].err, rip.loc[(M3, 0.15)].succ, rip.loc[(M3, 0.15)].err, rip.loc[(M3, 0.03)].succ - rip.loc[(M3, 0.15)].succ,
            noi.loc[(M3, 40.0)].succ, min(nz2[:4]), max(nz2[:4]), noi.loc[(M3, 60.0)].succ, noi.loc[(M2, 60.0)].succ))
    b.fig("stress", "Success at 0.25 m/s versus latency, ripple and image noise.", 11.0)
    b.h2("9.5 What happens during one grasp")
    b.p("{F:trace} compares the same 8-object trial for M2 and M3 at 0.25 m/s. Panel (a) shows the tip-to-object distance in the 0.2 s before contact. Both methods start about 50 mm apart, "
        "because the arm waits ahead of the cube; M3 converges smoothly to a few millimetres while the final distance of M2 depends on the accumulated prediction error. Panel (b) shows the error at contact for each pick: "
        "M2 varied between %.1f and %.1f mm and %d of its %d contacts were outside the 10 mm capture radius (misses that needed a second attempt), whereas M3 stayed between %.1f and %.1f mm with %d misses. "
        "{F:snap} shows three moments of an M3 trial rendered from the 3-D scene." % (min(g2), max(g2), sum(e >= 10 for e in g2), len(g2), min(g3), max(g3), sum(e >= 10 for e in g3)))
    b.fig("trace", "Same trial for M2 and M3: approach and error at contact.", 11.0)
    b.fig("snap", "Snapshots of one trial of the proposed method.", 12.0)
    b.h2("9.6 Ablation study")
    b.p("{T:abl} and {F:ablation} isolate the components in a hard condition (0.30 m/s, 6 %% ripple, 160 ms latency). Starting from the plain predictive method M2 (%.1f %%), adding the Kalman filter with encoder fusion "
        "and the visual-servo descent together gives the full M3 (%.1f %%). Removing the visual servo from M3 costs %.1f points of success and raises the grasp error from %.1f mm to %.1f mm; "
        "removing the Kalman filter costs %.1f points and raises the error to %.1f mm. Both components therefore contribute, and the visual servo contributes more. The scheduling rule makes no difference (%.1f %% with and without). "
        "With 10 seeds the standard deviations are 6 to 12 points, so the Kalman difference (Welch t-test p = 0.20) could not be separated from the run-to-run scatter. The confirmation study below, with 30 new seeds, removes this doubt. The ordering M2 < no servo < no Kalman < full M3 is clear, and full M3 also has the lowest grasp error of all variants (%.2f mm)." % (
            ab.loc[M2].succ, ab.loc[M3].succ, ab.loc[M3].succ - ab.loc["M3 w/o visual servo"].succ, ab.loc[M3].err, ab.loc["M3 w/o visual servo"].err,
            ab.loc[M3].succ - ab.loc["M3 w/o Kalman (LS)"].succ, ab.loc["M3 w/o Kalman (LS)"].err, ab.loc["M3 w/o SCT scheduling"].succ, ab.loc[M3].err))
    rows = [["Variant", "Success %", "Std (pts)", "Throughput obj/min", "Grasp error mm"]]
    for k, lab in ((M2, "M2 Predictive"), ("M3 w/o visual servo", "M3 without visual servo"), ("M3 w/o Kalman (LS)", "M3 without Kalman (line fit)"),
                   ("M3 w/o SCT scheduling", "M3 without SCT (FIFO)"), (M3, "M3 Proposed (full)")):
        rows.append([lab, "%.1f" % ab.loc[k].succ, "%.1f" % ab.loc[k].sd, "%.1f" % ab.loc[k].thr, "%.2f" % ab.loc[k].err])
    b.table("abl", "Ablation at 0.30 m/s, 6 % ripple, 160 ms latency (10 runs, 15 objects each).", rows, (5.6, 2.2, 2.2, 3.0, 2.5), size=9.5)
    b.fig("ablation", "Ablation: (a) sorting success and (b) grasp error of the variants.", 9.5)
    b.p("**Statistical confirmation.** To make the claims independent of the first 10 seeds, the key comparisons were repeated with 30 new seeds (330 runs, 15 parts each; {T:conf}). "
        "Every component of the proposed method gives a significant gain (Welch t-test of each variant against the full method): removing the Kalman filter costs %.1f points (p %s), removing the visual servo costs %.1f points (p %s) and "
        "the plain predictive method M2 is %.1f points lower (p %s). The scheduling rule remains without effect (identical results). {T:conf2} compares M3 with M2 under other conditions: the advantage is significant at 0.25 m/s and 0.30 m/s and very large at 15 %% belt-speed ripple." % (
            cab[M3].succ.mean() - cab["M3 w/o Kalman (LS)"].succ.mean(), pfmt(welch(cab[M3].succ, cab["M3 w/o Kalman (LS)"].succ)),
            cab[M3].succ.mean() - cab["M3 w/o visual servo"].succ.mean(), pfmt(welch(cab[M3].succ, cab["M3 w/o visual servo"].succ)),
            cab[M3].succ.mean() - cab[M2].succ.mean(), pfmt(welch(cab[M3].succ, cab[M2].succ))))
    rows = [["Variant (0.30 m/s, 6 % ripple, 160 ms)", "Success %", "Std (pts)", "Grasp error mm", "p vs full M3"]]
    for k, lab in ((M2, "M2 Predictive"), ("M3 w/o visual servo", "M3 without visual servo"), ("M3 w/o Kalman (LS)", "M3 without Kalman (line fit)"),
                   ("M3 w/o SCT scheduling", "M3 without SCT (FIFO)"), (M3, "M3 Proposed (full)")):
        rows.append([lab, "%.1f" % cab[k].succ.mean(), "%.1f" % cab[k].succ.std(), "%.2f" % cab[k].grasp_err_mm.mean(), "-" if k == M3 else pfmt(welch(cab[M3].succ, cab[k].succ))])
    b.table("conf", "Confirmation study of the ablation with 30 new seeds (450 parts per variant).", rows, (6.4, 2.2, 2.0, 2.7, 2.2), size=9.5)
    rows = [["Condition (30 seeds)", "M3 success %", "M2 success %", "Difference (pts)", "p"]]
    for lab, tg, key, val in (("0.25 m/s", "confirm_speed", "v", 0.25), ("0.30 m/s", "confirm_speed", "v", 0.30), ("0.25 m/s, 15 % ripple", "confirm_ripple", "belt_ripple", 0.15)):
        x_ = cf[(cf.tag == tg) & (cf[key] == val)]
        a_s, b_s = x_[x_.method == M3].succ, x_[x_.method == M2].succ
        rows.append([lab, "%.1f" % a_s.mean(), "%.1f" % b_s.mean(), "%.1f" % (a_s.mean() - b_s.mean()), pfmt(welch(a_s, b_s))])
    b.table("conf2", "M3 versus M2 in the confirmation study (30 new seeds).", rows, (5.2, 2.5, 2.5, 3.0, 2.3), size=9.5)
    b.h2("9.7 Cross-platform validation in CoppeliaSim")
    cs_speed = cs_[cs_.tag == "speed"].groupby(["method", "v"]).agg(succ=("succ", "mean"), thr=("throughput", "mean"), err=("grasp_err_mm", "mean"))
    n_cs_obj = int(cs_.n.sum())
    n_cs_ok = int(cs_.correct.sum())
    cs_fa, cs_fr, cs_def = int(cs_.false_accept.sum()), int(cs_.false_reject.sum()), int(cs_.defects.sum())
    cs_mis = int((cs_.correct != cs_.correct_logic).sum())
    csab = {m_: x_ for m_, x_ in cs_[cs_.tag == "ablation"].groupby("method")}
    b.p("The complete software was also run in CoppeliaSim, the assigned platform, with 105 runs of 12 parts each (%d parts, of which %d cylinders). {T:cs} compares it with the MuJoCo study. "
        "A part counts only if it rests in its correct bin in the CoppeliaSim physics; in %d of %d runs this check agreed with the planner log. "
        "The ranking of the methods is the same on both engines: M1 (stop-and-go) sorts everything but slowly, and M3 is at least as successful as M2 in %d of 4 speeds, with %.1f %% against %.1f %% at 0.30 m/s. "
        "The grasp error of M3 is %.1f to %.1f mm in CoppeliaSim and %.1f to %.1f mm in MuJoCo; it is larger in CoppeliaSim because the joint loop is evaluated every 5 ms and the Bullet contacts are softer. "
        "Inspection in CoppeliaSim: %d false accepts and %d false rejects among %d cylinders and %d cubes." % (
            n_cs_obj, cs_def, len(cs_) - cs_mis, len(cs_), cs_rank, cs_speed.loc[(M3, 0.30)].succ, cs_speed.loc[(M2, 0.30)].succ,
            cs_speed.loc[M3].err.min(), cs_speed.loc[M3].err.max(), sp.loc[M3].err.min(), sp.loc[M3].err.max(), cs_fa, cs_fr, cs_def, n_cs_obj - cs_def))
    rows = [["Method", "v (m/s)", "CoppeliaSim succ. %", "MuJoCo succ. %", "CoppeliaSim obj/min", "MuJoCo obj/min", "CS err mm", "MuJoCo err mm"]]
    for m_ in (M1, M2, M3):
        for v_ in (0.10, 0.20, 0.25, 0.30):
            rows.append([m_.split(" ", 1)[0], "%.2f" % v_, "%.1f" % cs_speed.loc[(m_, v_)].succ, "%.1f" % sp.loc[(m_, v_)].succ,
                         "%.1f" % cs_speed.loc[(m_, v_)].thr, "%.1f" % sp.loc[(m_, v_)].thr, "%.1f" % cs_speed.loc[(m_, v_)].err, "%.1f" % sp.loc[(m_, v_)].err])
    b.table("cs", "Speed study in CoppeliaSim (5 seeds, 12 parts) and in MuJoCo (10 seeds, 15 parts).", rows, (1.5, 1.4, 2.2, 2.0, 2.3, 2.0, 1.6, 1.9), size=8.5, align_left=False)
    def _cs(tag, mm, **kw):
        x_ = cs_[(cs_.tag == tag) & (cs_.method == mm)]
        return x_.succ.mean()

    def _mj(tag, mm, key, val):
        return df[(df.tag == tag) & (df.method == mm) & (df[key] == val)].succ.mean()
    b.p("The stress tests and the ablation give the same picture on both engines. At 15 %% belt-speed ripple M3 sorts %.1f %% (CoppeliaSim) and %.1f %% (MuJoCo) of the parts, M2 only %.1f %% and %.1f %%. "
        "With 320 ms vision latency M3 reaches %.1f %% and %.1f %%, M2 %.1f %% and %.1f %%. In the ablation at 0.30 m/s the success rises step by step from M2 (%.1f %% / %.1f %%) through M3 without visual servo (%.1f %% / %.1f %%) "
        "and M3 without Kalman filter (%.1f %% / %.1f %%) to the full method M3 (%.1f %% / %.1f %%). With only 5 seeds per point the CoppeliaSim differences are less precise than the MuJoCo results, which is why the statistics of Section 9.6 were done in MuJoCo." % (
            _cs("ripple", M3), _mj("ripple", M3, "belt_ripple", 0.15), _cs("ripple", M2), _mj("ripple", M2, "belt_ripple", 0.15),
            _cs("latency", M3), _mj("latency", M3, "latency", 0.32), _cs("latency", M2), _mj("latency", M2, "latency", 0.32),
            csab[M2].succ.mean(), ab.loc[M2].succ, csab["M3 w/o visual servo"].succ.mean(), ab.loc["M3 w/o visual servo"].succ,
            csab["M3 w/o Kalman (LS)"].succ.mean(), ab.loc["M3 w/o Kalman (LS)"].succ, csab[M3].succ.mean(), ab.loc[M3].succ))
    b.h2("9.8 Discussion: what worked and what did not")
    b.bullets([
        "**Worked:** predicting the intercept, correcting the grasp point during the descent and fusing the encoder. These keep the success high (M3: %.1f %% at 0.25 m/s) although the camera is 80 ms late, the belt speed fluctuates and the arm hides parts of the scene." % sp.loc[(M3, 0.25)].succ,
        "**Worked:** the minimum-jerk plus feedforward motion: it is smooth, respects the limits and is tracked within %.1f mrad." % trk["PID+FF"],
        "**Did not work:** the SCT scheduling rule (no gain) and chasing the current object position during the descent (our first version, which failed completely above 0.25 m/s because the arm arrives ahead of the cube by about 50 mm).",
        "**Limit:** above about 0.25 m/s the single arm is saturated at roughly %.0f objects per minute; a second arm or faster motion limits would be needed. The camera is orthographic, the conveyor is kinematic and the suction is an attach constraint, so absolute numbers will differ on a real robot." % sp.loc[M3].thr.max(),
    ])

    # ---------------------------------------------------------------- 10
    b.h1("10 Contribution")
    b.p("Each member owns one technical module. All claims below can be traced to the files of the GitHub repository (Appendix) and to the corresponding part of the video. "
        "The member who built a module presents it in the video and slides. {T:contrib} gives the contribution of each member.")
    contrib = [
        ("Built the SCARA kinematics in arm.py (forward and inverse kinematics with joint limits and elbow choice, Jacobian, D-H matrices, minimum-jerk generator and the duration formula) and the arm state machine and the visual-servo descent in sim.py (_arm, _grasp, _release). Also built the CoppeliaSim cell (cs_world.py: scene, arm, vision sensor, bins, suction attach; simcs.py: closed loop) and the MuJoCo cell in mj_world.py (arm, belt, bins, free-body cubes, orthographic camera, suction weld) "
         "and the closed loop in sim3d.py. Produced Fig. 1, 3, 10 and 11.", "Abstract, 1, 2, 12, Appendix"),
        ("Built the camera and detector in vision.py (rendering of rotated cubes and cylinders, noise and lighting, HSV thresholds, morphological filtering, connected components, contour shape classifier, homography calibration) "
         "and the vision study (extra_experiments.py: vision_study). Produced Fig. 5 and 7.", "3, 4, 8"),
        ("Built the tracker in tracker.py (nearest-neighbour association, Kalman filter with encoder pseudo-measurement and latency extrapolation, line-fit baseline) and the intercept planner in sim.py (_intercept, _plan, SCT/FIFO scheduling). "
         "Ran the latency and ablation studies. Produced Fig. 4 and 12.", "5, 6"),
        ("Built the joint plant and PID controller with feedforward in arm.py (JointPlant), derived the gains, ran the step-response and tracking tests (extra_experiments.py: step_response, traj_tracking) and tuned wn, zeta and the integrator clamp. "
         "Wrote the torque-level controller JointPlantMJ in sim3d.py and the joint-callback PID + feedforward script of the CoppeliaSim cell (cs_world.py), tested in extra_cs.py. Produced Fig. 6.", "7, 9.1-9.2, References"),
        ("Built the experiment harness (experiments.py, parallel runs and CSV logging), the stop-and-go baseline logic in sim.py (belt stop/go and static position estimate), make_figures.py, experiments_cs.py (8 parallel CoppeliaSim instances), cs_launch.py, demo_cs.py and render_cs.py (live CoppeliaSim demo, stills, video), demo.py and render_3d.py (MuJoCo), the unit tests and the README. "
         "Produced Fig. 2, 8, 9 and the result tables.", "9.3-9.8"),
    ]
    rows = [["S. No.", "Member name and SAP ID", "Technical contribution (modules / algorithms / code files / experiments)", "Report sections written + signature"]]
    for i, ((n, sid), (txt, secs)) in enumerate(zip(MEMBERS, contrib), 1):
        rows.append([str(i), "%s (%s)" % (n, sid), txt, "Sections %s\n\nSign: ________" % secs])
    b.table("contrib", "Technical contribution of each member and report sections written.", rows, (1.1, 3.3, 8.2, 2.9), size=8.5, keep=False)

    # ---------------------------------------------------------------- 11
    b.h1("11 Co-learning Reflection")
    p = d.add_paragraph()
    self_note = p.add_run("[Each member must replace the draft below with their own words before printing.]")
    highlight(self_note)
    refl = [
        ("Aditya Khatri", "On my own I learned how inverse kinematics, the Jacobian and D-H matrices turn into code, and why the elbow solution closest to the present pose avoids joint flips. "
         "Debugging my first visual servo showed me that the arm arrives ahead of a moving part and must aim at the predicted contact point. "
         "From my team-mates I learned how the Kalman filter and the PID gains shape my trajectories. I taught the team how minimum-jerk durations follow from speed and acceleration limits."),
        ("Simran Jeet Gill", "I learned HSV colour segmentation and homography calibration from the OpenCV documentation and by testing thresholds on rendered images, and I designed the contour feature that tells cylinders from cubes. "
         "I saw that shape fails before colour when image noise grows. From my team-mates I learned how detections feed the tracker and why latency matters. I taught the others how pixels become world coordinates."),
        ("Divya Israni", "I studied the Kalman filter from the tracking literature and implemented a constant-velocity model with an encoder pseudo-measurement. "
         "I learned that a filter only helps when latency is handled, which led to the extrapolation step. From my team-mates I learned the controller and vision details that feed my tracker. "
         "I taught the others how to read the ablation tables and why the scheduling rule gave no benefit."),
        ("Amritanshi Singh", "I learned to design a PID controller from a desired natural frequency and damping and to add model feedforward. "
         "Testing on the arm showed me that a P-only loop oscillates, that PD is smooth, and that the pose-dependent inertia of a coupled arm limits what feedforward can do. "
         "From my team-mates I learned how my controller is used in the visual-servo descent. I taught the team how to read the step and tracking plots."),
        ("Anchit Saxena", "I learned to compare methods fairly: same seeds, many runs, error bars and significance tests. "
         "Counting success from the physics state instead of the planner log showed me that parts pile up in a bin, which we fixed with a release grid. "
         "From my team-mates I learned the algorithms I evaluated. I taught the others how to run the scripts, rebuild the figures and start the live demo."),
    ]
    for n, t in refl:
        par = d.add_paragraph(); par.paragraph_format.space_after = Pt(2)
        par.add_run(n + ": ").bold = True
        par.add_run(t)

    # ---------------------------------------------------------------- 12
    b.h1("12 Conclusion and Future Scope")
    b.p("A complete vision-guided colour-and-shape sorting cell for quality inspection was designed and simulated in CoppeliaSim (the assigned platform) and in the MuJoCo physics engine: an OpenCV detector on rendered camera images, a latency-compensated Kalman tracker with encoder fusion, an intercept planner, "
        "minimum-jerk motion, joint PID with feedforward and a predictive visual-servo descent. Over 560 MuJoCo runs, in which a part counts only if it physically rests in the right bin, the proposed method sorted %.1f %% of the cubes at 0.20 m/s and %.1f %% at 0.25 m/s, "
        "kept the mean grasp error between %.1f and %.1f mm (the perception floor of this camera is about %.1f mm), stayed robust at 15 %% belt-speed ripple (%.1f %% success against %.1f %% for the plain predictive method) and delivered %.1f times the throughput of a stop-and-go conveyor at 0.25 m/s. "
        "Inspection errors were extremely rare: one false accept in more than 3000 cylinders and no false reject at image noise up to 40. The same software reproduced the ranking of the methods in 105 CoppeliaSim runs (Section 9.7). The ablation showed that the visual-servo descent matters most, the Kalman filter helps significantly too (30-seed confirmation study, Section 9.6), and the scheduling rule did not help." % (
            sp.loc[(M3, 0.2)].succ, sp.loc[(M3, 0.25)].succ, sp.loc[M3].err.min(), sp.loc[M3].err.max(), sp.loc[M1].err.mean(),
            rip.loc[(M3, 0.15)].succ, rip.loc[(M2, 0.15)].succ, sp.loc[(M3, 0.25)].thr / sp.loc[(M1, 0.25)].thr))
    b.p("**Industrial use.** The same structure applies to quality inspection and sorting in manufacturing (the application domain of this project), parcel and component sorting, food grading and recycling lines. The study also gives a design rule: the capture radius of the gripper "
        "and the ripple of the belt decide whether open-loop prediction is enough or whether a feedback correction in the descent is needed.")
    b.p("**Transfer to a real robot.** The modules map directly onto hardware: a real camera with a hardware trigger replaces the renderer, the encoder feeds the filter, the planner and the visual servo run on the robot controller "
        "and the joint PID is replaced by the drive's servo loop. Remaining work includes hand-eye calibration, lens distortion, gripper vacuum delay and collision checking. "
        "Future extensions are a second arm to raise the capacity, learning-based detection for unknown objects, orientation-aware grasping with a gripper, an unscented filter for varying belt speed, a ROS integration and a hardware-in-the-loop test on a real SCARA arm.")

    # ---------------------------------------------------------------- 13
    b.h1("13 References")
    ref_par = d.add_paragraph()   # placeholder, replaced below
    refs_anchor = ref_par

    # ---------------------------------------------------------------- Appendix
    b.h1("Appendix: Links and Logs")
    b.h2("A.1 Project links")
    par = d.add_paragraph(); par.paragraph_format.line_spacing = 1.3
    par.add_run("GitHub repository: ").bold = True
    r = par.add_run(GH); highlight(r)
    par = d.add_paragraph(); par.paragraph_format.line_spacing = 1.3
    par.add_run("YouTube video (unlisted): ").bold = True
    r = par.add_run(YT); highlight(r)
    b.h2("A.2 Meeting log")
    b.p("{T:log} records the meetings of the group.")
    rows = [["No.", "Date", "Duration", "Attendees", "Decisions"]]
    for dec in ("Group formed, lead chosen, modules divided, simulator decided (CoppeliaSim + Python/OpenCV)", "Architecture and parameters agreed; GitHub repository created",
                "Results discussed; novelty points fixed; report writing split",
                "Report and slides reviewed; video rehearsal"):
        rows.append([str(len(rows)), "__/__/2026", "__ min", "All 5", dec])
    b.table("log", "Meeting log of the group.", rows, (1.0, 2.5, 1.9, 1.8, 8.3), size=9.5)
    b.h2("A.3 Peer-assessment sheets")
    b.p("One sheet per member ({T:peer}): each member scores every team-mate from 1 to 5 and signs the own sheet in the printed copy.")
    rows = [["No.", "Team-mate name and SAP ID", "Contribution (1-5)", "Reliability / on time (1-5)", "Helped others learn (1-5)", "Quality of work (1-5)"]]
    for mi, (mn, msid) in enumerate(MEMBERS):
        rows.append(["", "**Sheet %d, filled by: %s (%s)**" % (mi + 1, mn, msid), "", "", "", ""])
        k = 0
        for tj, (tn, tsid) in enumerate(MEMBERS):
            if tj == mi:
                continue
            k += 1
            rows.append([str(k), "%s (%s)" % (tn, tsid), "", "", "", ""])
    b.table("peer", "Peer-assessment sheets (one per member).", rows, (1.1, 5.6, 2.2, 2.4, 2.3, 1.9), size=8, keep=False)

    # ---- references (need citation order) ----
    order = b.cites
    missing = [k for k in REFS if k not in order]
    assert not missing, "uncited references: %s" % missing
    first = True
    anchor = refs_anchor
    prev = anchor
    for i, k in enumerate(order, 1):
        par = d.add_paragraph()
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.left_indent = Cm(1.0)
        par.paragraph_format.first_line_indent = Cm(-1.0)
        par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.line_spacing = 1.15
        par.add_run("[%d]\t" % i)
        b._runs(par, REFS[k])
        prev._p.addnext(par._p)
        prev = par
    anchor._p.getparent().remove(anchor._p)

    # superscript/subscript markup cleanup done in post-processing (see fix_sub)
    fix_sub(d)
    d.save(out_path)
    return b


def fix_sub(doc):
    """Convert <sub>..</sub>, <super>..</super> and html entities that were written inline into real sub/superscripts."""
    import html

    def process(par):
        txt = "".join(r.text for r in par.runs)
        if "<sub>" not in txt and "<super>" not in txt and "&#" not in txt:
            return
        for r in list(par.runs):
            if not any(t in r.text for t in ("<sub>", "<super>", "&#")):
                continue
            pieces = re.split(r"(<sub>.*?</sub>|<super>.*?</super>)", r.text)
            if len(pieces) == 1:
                r.text = html.unescape(r.text)
                continue
            new = []
            for pc in pieces:
                if not pc:
                    continue
                if pc.startswith("<sub>"):
                    new.append((html.unescape(pc[5:-6]), "sub"))
                elif pc.startswith("<super>"):
                    new.append((html.unescape(pc[7:-8]), "sup"))
                else:
                    new.append((html.unescape(pc), None))
            anchor = r._r
            for text, kind in new:
                nr = copy_run(r, text)
                if kind == "sub":
                    nr.font.subscript = True
                elif kind == "sup":
                    nr.font.superscript = True
                anchor.addnext(nr._r)
                anchor = nr._r
            r._r.getparent().remove(r._r)

    def copy_run(r, text):
        import copy
        from docx.text.run import Run
        el = copy.deepcopy(r._r)
        for t in el.findall(qn("w:t")):
            el.remove(t)
        nr = Run(el, r._parent)
        nr.text = text
        return nr

    def walk(container):
        for par in container.paragraphs:
            process(par)
        for tb in container.tables:
            for row in tb.rows:
                for cell in row.cells:
                    walk(cell)
    walk(doc)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "AISR_G10_Project_Report.docx")
    pages = {}
    pj = os.path.join(HERE, "pages.json")
    if os.path.exists(pj):
        pages = json.load(open(pj))
    build(out, pages)
    print("saved", out)
