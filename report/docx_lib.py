"""Small helper layer over python-docx for the AISR report format."""
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"


def set_font(style_or_run, size=None, bold=None, italic=None, name=FONT):
    f = style_or_run.font
    f.name = name
    el = style_or_run.element
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el.rPr
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    f.color.rgb = RGBColor(0, 0, 0)


def fmt(style, align=None, before=0, after=0, line=None, keep_next=False):
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line
    pf.keep_with_next = keep_next


def add_field(par, instr, placeholder=""):
    r = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    r2 = par.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " %s " % instr; r2._r.append(it)
    r3 = par.add_run()
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r3._r.append(s)
    r4 = par.add_run(placeholder)
    r5 = par.add_run()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r5._r.append(e)
    return r4


def page_border(section, space=20):
    sp = section._sectPr
    for old in sp.findall(qn("w:pgBorders")):
        sp.remove(old)
    pb = OxmlElement("w:pgBorders")
    pb.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:space"), str(space)); e.set(qn("w:color"), "000000")
        pb.append(e)
    sp.find(qn("w:pgMar")).addnext(pb)
    return pb


def page_numbering(section, fmt_="decimal", start=1):
    sp = section._sectPr
    for old in sp.findall(qn("w:pgNumType")):
        sp.remove(old)
    pn = OxmlElement("w:pgNumType")
    pn.set(qn("w:fmt"), fmt_); pn.set(qn("w:start"), str(start))
    sp.find(qn("w:pgBorders")).addnext(pn)


def footer_pagenum(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    r = add_field(p, "PAGE", "1")
    return p


def footer_empty(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)


def shade(cell, color="E7E6E6"):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def highlight(run, color="yellow"):
    from docx.enum.text import WD_COLOR_INDEX
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
